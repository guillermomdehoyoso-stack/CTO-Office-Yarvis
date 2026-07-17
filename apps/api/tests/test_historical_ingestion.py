from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from yarvis_api.schemas.historical_evidence import (
    HistoricalEvidenceRecord,
    HistoricalWorkbookInspection,
    ProfitabilityDuplicateResult,
    TransactionHistoricalAggregate,
)
from yarvis_api.services.historical_ingestion import (
    PROFITABILITY_REPORT,
    TRANSACTION_DETAIL_REPORT,
    aggregate_transaction_evidence,
    detect_profitability_duplicates,
    normalize_decimal,
    normalize_header,
    normalize_historical_workbook,
    normalize_identifier,
    normalize_reporting_month,
)


def _write_workbook(path: Path, collision: bool = False, missing: bool = False) -> None:
    workbook = Workbook()
    transactions = workbook.active
    transactions.title = "BD Detalle por Transaccion"
    headers = [" Store ID ", "Client ID", "Tipo de Transacción", "Transacciones", "Volumen", "Tasa de Venta", "Tasa de Compra", "Diferencial de Tasas", "Mes", "Producto", "Tarjeta", "Banco Adquirente", "Banco Emisor", "Procesador", "Unmapped Column"]
    if collision:
        headers[1] = " store   id "
    if missing:
        headers.remove("Volumen")
    transactions.append(headers)
    transactions.append(["001", "C-1", "sale", 2, "1,000.50", "1", "2", "0.5", "2025-01", "P1", "V", "A", "I", "X", "unresolved"])
    transactions.append(["001", "C-1", "sale", 3, 20, 1, 2, Decimal("0.5"), date(2025, 1, 31), "P2", "V", "A", "I", "X", "unresolved"])
    transactions.append(["001", "C-1", "refund", 1, 5, 1, 2, 1, "2025-02-02", "P1", "M", "B", "J", "Y", "unresolved"])
    transactions.append([1.5, "C-1", "sale", "bad", 1, 1, 2, 1, "bad-month", "P1", "V", "A", "I", "X", "unresolved"])
    profitability = workbook.create_sheet("BD Rentabilidad")
    p_headers = ["Store ID", "Client ID", "Branch", "Volumen", "Renta Pagada", "Total Comisiones (sin Activación ni Referenciador)", "Rentabilidad", "Mes", "Store Name"]
    profitability.append(p_headers)
    profitability.append(["002", "C-2", "branch", 10, 1, 2, 3, "2025-01", "synthetic"])
    profitability.append(["002", "C-2", "branch", 10, 1, 2, 3, "2025-01", "synthetic"])
    profitability.append(["003", "C-3", "branch", 11, 1, 2, 3, "2025-01", "synthetic"])
    profitability.append(["003", "C-3", "branch", 12, 1, 2, 3, "2025-01", "synthetic"])
    profitability.append([None, "C-4", "branch", 1, 1, 2, 3, "2025-01", "synthetic"])
    workbook.create_sheet("Ignored Summary").append(["summary"])
    workbook.save(path)


def _records(tmp_path: Path):
    path = tmp_path / "synthetic_historical.xlsx"
    _write_workbook(path)
    return normalize_historical_workbook(path, datetime(2025, 2, 1, tzinfo=timezone.utc))


def test_header_and_value_normalization():
    assert normalize_header("  RAZÓN   Social ") == "razon social"
    assert normalize_identifier("001")[0] == "001"
    assert normalize_identifier(12.0)[0] == "12"
    assert normalize_identifier(1.5)[0] is None
    assert normalize_identifier(None)[0] is None
    assert normalize_reporting_month(date(2025, 1, 2))[0] == "2025-01"
    assert normalize_reporting_month(datetime(2025, 2, 2))[0] == "2025-02"
    assert normalize_reporting_month("2025-03-10")[0] == "2025-03"
    assert normalize_reporting_month("2025-13")[0] is None
    assert normalize_decimal(" $1,000.50 ")[0] == Decimal("1000.50")
    assert normalize_decimal("5%")[0] is None
    assert normalize_decimal("not-number")[0] is None


def test_typed_inspection_normalization_and_provenance(tmp_path):
    inspection, records = _records(tmp_path)
    assert isinstance(inspection, HistoricalWorkbookInspection)
    assert inspection.supported_sheets == ["BD Detalle por Transaccion", "BD Rentabilidad"]
    assert inspection.ignored_sheets == ["Ignored Summary"]
    assert all(isinstance(record, HistoricalEvidenceRecord) for record in records)
    transaction = next(record for record in records if record.evidence_type == TRANSACTION_DETAIL_REPORT)
    assert transaction.temporal_class == "historical"
    assert transaction.source_row_reference == "row:2"
    assert transaction.original_headers["store_id"] == " Store ID "
    assert transaction.unresolved_source_fields["Unmapped Column"] == "unresolved"
    invalid = next(record for record in records if record.source_sheet == "BD Detalle por Transaccion" and record.source_row_reference == "row:5")
    assert invalid.identity_review_required
    assert {error.error_code for error in invalid.validation_errors} >= {"invalid_store_id", "invalid_reporting_period", "invalid_numeric_value", "missing_store_id", "missing_reporting_period"}


def test_collided_and_partial_sheets_are_not_promotable(tmp_path):
    collision = tmp_path / "collision.xlsx"
    _write_workbook(collision, collision=True)
    inspected, records = normalize_historical_workbook(collision, datetime.now(timezone.utc))
    transaction = next(item for item in inspected.sheet_results if item.source_sheet == "BD Detalle por Transaccion")
    assert not transaction.promotable and "normalized_header_collision" in transaction.structural_conflicts
    assert not [record for record in records if record.evidence_type == TRANSACTION_DETAIL_REPORT]
    partial = tmp_path / "partial.xlsx"
    _write_workbook(partial, missing=True)
    inspected, _ = normalize_historical_workbook(partial, datetime.now(timezone.utc))
    assert "BD Detalle por Transaccion" in inspected.unresolved_sheets


def test_transaction_aggregation_is_typed_month_scoped_and_decimal(tmp_path):
    _, records = _records(tmp_path)
    aggregates, excluded = aggregate_transaction_evidence(records)
    assert all(isinstance(item, TransactionHistoricalAggregate) for item in aggregates)
    jan = next(item for item in aggregates if item.store_id == "001" and item.reporting_period == "2025-01")
    assert jan.historical_transaction_count == Decimal("5")
    assert jan.historical_transaction_volume == Decimal("1020.50")
    assert jan.historical_rate_margin is None
    assert jan.rate_margin_observations == [Decimal("0.5")]
    assert jan.products_observed == ["P1", "P2"]
    assert any(record.source_row_reference == "row:5" for record in excluded)
    assert len(aggregates) == 2


def test_profitability_duplicates_are_explicit_and_never_last_write_wins(tmp_path):
    _, records = _records(tmp_path)
    result = detect_profitability_duplicates(records)
    assert isinstance(result, ProfitabilityDuplicateResult)
    assert [(item.store_id, item.record_count) for item in result.exact_duplicate_groups] == [("002", 2)]
    assert [(item.store_id, item.record_count) for item in result.conflicting_duplicate_groups] == [("003", 2)]
    assert [item.store_id for item in result.accepted_summaries] == ["002"]
    assert len(result.excluded_records) == 1
