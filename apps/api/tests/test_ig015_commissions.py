"""IG-015 acceptance against real isolated PostgreSQL; synthetic sources only."""

import hashlib
import zipfile
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import func, select
from test_netpay_operational_data import headers, setup_data, upload

from yarvis_api.main import app
from yarvis_api.models.netpay_operational_data import OperationalDataBatch, OperationalDataRow, StoreProfitabilityFact
from yarvis_api.models.principal import PrincipalMembership
from yarvis_api.services import netpay_commission_settlement as settlement_service
from yarvis_api.services.netpay_commerce_attention import NetpayCommerceAttentionService
from yarvis_api.services.netpay_commission_settlement import COMPONENTS, normalize_ooxml_package, parse_settlement

client = TestClient(app)


def core(ap="10", aq="5", ar="85", benefit="700", volume="1000", period="2026-09", key="core", reorder=False):
    labels = ["Store ID", "Mes", "Volumen", "Rentabilidad", *COMPONENTS.values()]
    values = ["SYN-STORE-001", period, volume, benefit, ap, aq, ar]
    if reorder:
        labels.reverse()
        values.reverse()
    return upload(
        "synthetic.csv",
        (",".join(labels) + "\n" + ",".join(values) + "\n").encode(),
        "monthly_store_profitability",
        key,
    )


def accept(batch, key="accept", **extra):
    return client.post(
        f"/netpay/data/datasets/{batch['id']}/accept",
        headers=headers(key=key),
        json={"preview_token": batch["preview_token"], **extra},
    )


def metric(store, period="2026-09"):
    with app.state.yarvis.persistence.create_session() as db:
        fact = db.scalar(select(StoreProfitabilityFact).where(StoreProfitabilityFact.store_reference_id == store))
        from yarvis_api.models.netpay_master import NetpayStoreReference

        return NetpayCommerceAttentionService(db, fact.organization_id, period)._p2(db.get(NetpayStoreReference, store))


def workbook(movements=None, total="11", period="Septiembre 2026", repeat=False):
    w = Workbook()
    s = w.active
    s.title = "sheet1"
    s.append(["NOMBRE", None, "CORREO", None, "Total Comisiones", "Mes Corriente", "Mes Previo", "Mensaje"])
    s.append(
        [
            "Private synthetic",
            None,
            "private@example.invalid",
            None,
            total,
            "Octubre",
            "Septiembre",
            "Comisiones Septiembre 2026 para el pago del 15 de Octubre de 2026",
        ]
    )
    s.append(["NetPay"])
    s.append(["Comisiones Synthetic"])
    s.append(["Periodo " + period])
    s.append([])
    h = ["Entidad", "NOMBRE", "Mes", "Store", "Comercio", "CONCEPTO", "Total Comisiones", "CORREO"]
    s.append(h)
    for i, (entity, store, concept, amount) in enumerate(
        movements
        or [
            ("Asesor", "SYN-STORE-001", "COMISION", "10"),
            ("Distribuidor", "SYN-STORE-001", "BONO", "3"),
            ("Asociado 15", None, "DEDUCTIVA", "-2"),
        ]
    ):
        if repeat and i:
            s.append(h)
        s.append(
            [
                entity,
                "Private synthetic",
                "Septiembre",
                store,
                "Synthetic commerce",
                concept,
                amount,
                "private@example.invalid",
            ]
        )
    b = BytesIO()
    w.save(b)
    return b.getvalue()


def settlement(content=None, key="settlement", name="settlement.xlsx"):
    return upload(name, content or workbook(), "commission_settlement", key)


def rewrite_package(content, *, rename=lambda name: name, payload=lambda name, value: value, extras=()):
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(content)) as source, zipfile.ZipFile(output, "w") as target:
        for entry in source.infolist():
            target.writestr(rename(entry.filename), payload(entry.filename, source.read(entry)))
        for name, value in extras:
            target.writestr(name, value)
    return output.getvalue()


def noncanonical_relationship_package(content, *, omit_dimension=False):
    def rename(name):
        return name.replace("xl/_rels/workbook.xml.rels", "xl\\_rels\\workbook.xml.rels").replace(
            "_rels/.rels", "_rels\\.rels"
        )

    def payload(name, value):
        if omit_dimension and name.startswith("xl/worksheets/"):
            import re

            return re.sub(rb"<dimension[^>]*/>", b"", value)
        return value

    return rewrite_package(content, rename=rename, payload=payload)


def test_ooxml_compatibility_is_noop_for_canonical_and_deterministic_for_backslashes():
    canonical = workbook()
    assert normalize_ooxml_package(canonical) == canonical
    expected = parse_settlement(canonical)
    noncanonical = noncanonical_relationship_package(canonical, omit_dimension=True)
    first = parse_settlement(noncanonical)
    second = parse_settlement(noncanonical)
    assert first == second == expected


def test_ooxml_compatibility_tolerates_merged_nontransaction_cells():
    from openpyxl import load_workbook

    source = workbook()
    loaded = load_workbook(BytesIO(source))
    sheet = loaded.active
    sheet.merge_cells("B3:H3")
    output = BytesIO()
    loaded.save(output)
    loaded.close()
    movements, _, metadata = parse_settlement(noncanonical_relationship_package(output.getvalue(), omit_dimension=True))
    assert len(movements) == 3
    assert metadata["net_reported"] == "11"


def test_ooxml_compatibility_preserves_original_source_hash_and_rows():
    setup_data()
    original = noncanonical_relationship_package(
        workbook([("Asesor", "SYN-STORE-001", "COMISION", "5")] * 2, total="10")
    )
    response = settlement(original, "noncanonical")
    assert response.status_code == 201, response.text
    batch_id = UUID(response.json()["id"])
    with app.state.yarvis.persistence.create_session() as db:
        batch = db.get(OperationalDataBatch, batch_id)
        assert batch.source_hash == hashlib.sha256(original).hexdigest()
        assert db.scalar(select(func.count()).select_from(OperationalDataRow)) == 2


def test_settlement_accepts_parenthesized_source_filename():
    setup_data()
    response = settlement(key="parenthesized-name", name="GUILLERMO MARIO DE HOYOS OLIVERA (3).xlsx")
    assert response.status_code == 201, response.text


def test_ooxml_compatibility_rejects_canonical_collision_and_duplicate_entries():
    canonical = workbook()
    with zipfile.ZipFile(BytesIO(canonical)) as source:
        rels = source.read("xl/_rels/workbook.xml.rels")
        content_types = source.read("[Content_Types].xml")
    collision = rewrite_package(canonical, extras=(("xl\\_rels\\workbook.xml.rels", rels),))
    duplicate = rewrite_package(canonical, extras=(("[Content_Types].xml", content_types),))
    from fastapi import HTTPException

    with pytest.raises(HTTPException, match="workbook_entry_collision"):
        parse_settlement(collision)
    with pytest.raises(HTTPException, match="duplicate_workbook_entry"):
        parse_settlement(duplicate)

    case_collision = rewrite_package(canonical, extras=(("XL/workbook.xml", b"x"),))
    with pytest.raises(HTTPException, match="workbook_entry_collision"):
        parse_settlement(case_collision)


@pytest.mark.parametrize(
    "unsafe", ["/evil.xml", "C:\\evil.xml", "..\\evil.xml", "\\\\server\\evil.xml", "xl//evil.xml"]
)
def test_ooxml_compatibility_rejects_unsafe_paths(unsafe):
    from fastapi import HTTPException

    with pytest.raises(HTTPException, match="unsafe_workbook_path"):
        parse_settlement(rewrite_package(workbook(), extras=((unsafe, b"x"),)))


def test_ooxml_compatibility_rejects_zip_bomb_external_and_unsupported_type():
    from fastapi import HTTPException

    bomb_output = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(workbook())) as source,
        zipfile.ZipFile(bomb_output, "w", compression=zipfile.ZIP_DEFLATED) as target,
    ):
        for entry in source.infolist():
            target.writestr(entry.filename, source.read(entry))
        target.writestr("xl/media/bomb.bin", b"0" * 1_000_000)
    bomb = bomb_output.getvalue()
    with pytest.raises(HTTPException, match="workbook_compression_ratio_exceeded"):
        parse_settlement(bomb)

    def external(name, value):
        if name == "xl/_rels/workbook.xml.rels":
            return value.replace(b" Target=", b' TargetMode="External" Target=', 1)
        return value

    with pytest.raises(HTTPException, match="external_workbook_relationship"):
        parse_settlement(rewrite_package(workbook(), payload=external))

    def unsupported(name, value):
        if name == "[Content_Types].xml":
            return value.replace(
                b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
            )
        return value

    with pytest.raises(HTTPException, match="unsupported_workbook_type"):
        parse_settlement(rewrite_package(workbook(), payload=unsupported))


def test_ooxml_compatibility_rejects_entry_count_and_expanded_size_limits(monkeypatch):
    from fastapi import HTTPException

    canonical = workbook()
    monkeypatch.setattr(settlement_service, "MAX_PACKAGE_ENTRIES", 9)
    with pytest.raises(HTTPException, match="workbook_package_limits_exceeded"):
        parse_settlement(rewrite_package(canonical, extras=(("extra.xml", b"x"),)))

    monkeypatch.setattr(settlement_service, "MAX_PACKAGE_ENTRIES", 500)
    monkeypatch.setattr(settlement_service, "MAX_PACKAGE_ENTRY_BYTES", 3)
    with pytest.raises(HTTPException, match="workbook_package_limits_exceeded"):
        parse_settlement(canonical)

    monkeypatch.setattr(settlement_service, "MAX_PACKAGE_ENTRY_BYTES", 20_000_000)
    monkeypatch.setattr(settlement_service, "MAX_PACKAGE_EXPANDED_BYTES", 10)
    with pytest.raises(HTTPException, match="workbook_package_limits_exceeded"):
        parse_settlement(canonical)


def test_ooxml_compatibility_rejects_encrypted_entries():
    from fastapi import HTTPException

    encrypted = bytearray(workbook())
    for signature, offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        position = 0
        while True:
            position = encrypted.find(signature, position)
            if position < 0:
                break
            flag = int.from_bytes(encrypted[position + offset : position + offset + 2], "little") | 1
            encrypted[position + offset : position + offset + 2] = flag.to_bytes(2, "little")
            position += 4
    with pytest.raises(HTTPException, match="unsafe_workbook_entry"):
        parse_settlement(bytes(encrypted))


def test_ooxml_public_rejection_is_controlled_and_does_not_leak_paths():
    setup_data()
    malicious = rewrite_package(workbook(), extras=(("..\\private\\secret.xml", b"x"),))
    response = settlement(malicious, "unsafe-path")
    assert response.status_code == 422
    assert response.json() == {"detail": "unsafe_workbook_path"}
    assert "private" not in response.text and "secret" not in response.text


def test_ooxml_compatibility_truncated_and_invalid_xml_are_controlled():
    from fastapi import HTTPException

    with pytest.raises(HTTPException, match="invalid_settlement_workbook"):
        parse_settlement(workbook()[:100])

    def invalid_xml(name, value):
        return b"<invalid" if name == "xl/_rels/workbook.xml.rels" else value

    with pytest.raises(HTTPException, match="invalid_settlement_workbook"):
        parse_settlement(rewrite_package(workbook(), payload=invalid_xml))


@pytest.mark.parametrize(
    "ap,aq,ar,benefit,volume,expected,rate,reason",
    [
        ("10", "5", "85", "700", "1000", "100", "10.00", None),
        ("0", "0", "0", "700", "1000", "0", "0.00", None),
        ("0", "0", "-25", "700", "1000", "-25", "-2.50", None),
        ("10", "5", "85", "700", "0", "100", None, "volume_zero"),
        ("", "5", "85", "700", "1000", None, None, "commission_missing"),
        ("10", "", "85", "700", "1000", None, None, "commission_missing"),
        ("10", "5", "", "700", "1000", None, None, "commission_missing"),
        ("10", "5", "85", "700", "", "100", None, "volume_missing"),
        ("10", "5", "85", "", "1000", "100", "10.00", None),
        ("0.005", "0.005", "0", "999", "1", "0.010", "1.00", None),
    ],
)
def test_core_calculation(ap, aq, ar, benefit, volume, expected, rate, reason):
    store = setup_data()
    response = core(ap, aq, ar, benefit, volume, reorder=True)
    assert response.status_code == 201, response.text
    batch = response.json()
    assert accept(batch).status_code == 200
    p2 = metric(store)
    assert p2.distributor_commission_mxn == (Decimal(expected) if expected else None)
    assert p2.effective_commission_on_sales_percent == (Decimal(rate) if rate else None)
    assert p2.percentage_reason == reason
    assert p2.combined_benefit_mxn == (Decimal(benefit) if benefit else None)
    assert p2.metric_definition_version == "ap_aq_ar_v2"
    result = client.get(f"/netpay/data/datasets/{batch['id']}/results", headers=headers())
    assert result.status_code == 200, result.text


def test_legacy_preserved_and_reprocessing_blocked():
    store = setup_data()
    b = upload(
        "legacy.csv",
        b"Store ID,Mes,Volumen,Rentabilidad\nSYN-STORE-001,2026-09,1000,700\n",
        "monthly_store_profitability",
        "legacy",
    ).json()
    assert accept(b).status_code == 200
    p = metric(store)
    assert p.distributor_commission_mxn is None and p.combined_benefit_mxn == 700
    assert p.quality == "pending_reprocessing"
    with app.state.yarvis.persistence.create_session() as db:
        fact = db.scalar(select(StoreProfitabilityFact))
        fid = str(fact.id)
        assert fact.distributor_commission_exact == 700
    new = core().json()
    assert (
        accept(
            new, "correction", expected_predecessors={str(store): fid}, correction_reason="Synthetic correction"
        ).json()["detail"]
        == "legacy_reprocessing_requires_verified_source_workflow"
    )


def test_settlement_preserves_entities_rows_missing_store_and_replay():
    setup_data()
    content = workbook(repeat=True)
    response = settlement(content)
    assert response.status_code == 201, response.text
    b = response.json()
    assert b["row_counts"]["received"] == 3
    assert [r["source_row_number"] for r in b["rows"]] == [8, 10, 12]
    assert "private@" not in response.text and "Private synthetic" not in response.text
    assert accept(b).status_code == 200
    replay = settlement(content, "other-key", "renamed.xlsx")
    assert replay.status_code == 200 and replay.json()["id"] == b["id"]
    result = client.get(f"/netpay/data/datasets/{b['id']}/results", headers=headers()).json()
    assert [r["entity_original"] for r in result] == ["Asesor", "Distribuidor", "Asociado 15"]
    assert sum(Decimal(r["amount_signed"]) for r in result) == 11
    assert result[2]["store_id"] is None and result[2]["bank_status"] == "not_confirmed"
    assert client.get(f"/netpay/data/datasets/{b['id']}", headers=headers("data:foreign")).status_code == 404
    assert settlement(workbook(total="12"), "settlement").status_code == 409


def test_control_blocks_acceptance_and_bad_rows_are_not_dropped():
    setup_data()
    b = settlement(workbook(total="12")).json()
    assert accept(b).json()["detail"] == "settlement_control_difference"
    missing = workbook([("Asesor", "SYN-STORE-001", "COMISION", None)])
    assert settlement(missing, "missing").status_code == 422
    assert settlement(workbook(period="Octubre 2026"), "period").status_code == 422


def test_same_entity_identical_rows_are_distinct_and_concurrent_acceptance_once():
    setup_data()
    move = ("Asesor", "SYN-STORE-001", "COMISION", "5")
    b = settlement(workbook([move, move], total="10")).json()
    with ThreadPoolExecutor(2) as pool:
        results = list(pool.map(lambda key: accept(b, key), ["accept-a", "accept-b"]))
    assert all(r.status_code == 200 for r in results)
    with app.state.yarvis.persistence.create_session() as db:
        assert db.scalar(select(func.count()).select_from(OperationalDataRow)) == 2
        from yarvis_api.models.domain_event import DomainEvent

        assert (
            db.scalar(
                select(func.count())
                .select_from(DomainEvent)
                .where(DomainEvent.event_type == "netpay.operational_dataset.accepted")
            )
            == 1
        )


def test_settlement_correction_explicit_predecessor_and_reconciliation_permissions():
    setup_data()
    b = settlement().json()
    assert accept(b).status_code == 200
    newer = settlement(workbook(total="12", movements=[("Distribuidor", "UNKNOWN", "NEW", "12")]), "new").json()
    assert accept(newer, "no-reason").status_code == 409
    assert (
        accept(
            newer,
            "with-reason",
            expected_predecessors={"settlement": b["id"]},
            correction_reason="Synthetic replacement",
        ).status_code
        == 200
    )
    assert client.get("/netpay/commerce-attention/settlements", headers=headers()).status_code == 403
    with app.state.yarvis.persistence.create_session() as db:
        batch = db.get(OperationalDataBatch, UUID(b["id"]))
        member = db.scalar(
            select(PrincipalMembership).where(PrincipalMembership.organization_id == batch.organization_id)
        )
        member.role = "netpay_operations_operator"
        db.commit()
    read = client.get("/netpay/commerce-attention/settlements?include_details=true", headers=headers())
    assert read.status_code == 200, read.text
    item = read.json()[0]
    assert item["batch_id"] == newer["id"] and item["unexplained_difference"] is None
    assert item["bank_status"] == "not_confirmed"
    assert "store_unmatched" in item["movements"][0]["reasons"]
    assert client.get("/netpay/commerce-attention/settlements?reporting_period=2026-10", headers=headers()).json() == []


def test_duplicate_headers_rejected_and_as_does_not_change_numerator():
    setup_data()
    b = core().json()
    assert accept(b).status_code == 200
    malformed = b"Store ID,Mes,Rentabilidad,Rentabilidad\nSYN-STORE-001,2026-09,1,2\n"
    assert upload("duplicate.csv", malformed, "monthly_store_profitability", "dup").status_code == 422
    w = Workbook()
    s = w.active
    s.title = "BD Rentabilidad"
    s.append(["Store ID", "Mes", "Volumen", "Rentabilidad", *COMPONENTS.values()])
    s.append(["SYN-STORE-001", "2026-10", 1, 999, 0.005, 0.005, 0])
    out = BytesIO()
    w.save(out)
    new = upload("core.xlsx", out.getvalue(), "monthly_store_profitability", "xlsx")
    assert new.status_code == 201, new.text
    assert accept(new.json(), "xlsx-accept").status_code == 200


def test_exact_period_and_zero_sales_monetary_coverage():
    from datetime import datetime, timezone
    from uuid import uuid4

    from yarvis_api.schemas.netpay_commerce_attention import PortfolioContact, PortfolioP2Metrics, PortfolioRow

    def row(volume, commission, period="2026-09"):
        return PortfolioRow(
            row_key=str(uuid4()),
            client_id=uuid4(),
            entity_scope="client",
            commerce_name="Synthetic",
            updated_at=datetime.now(timezone.utc),
            contact=PortfolioContact(available=False),
            p2=PortfolioP2Metrics(
                reporting_period=period,
                monthly_sales_mxn=Decimal(volume),
                distributor_commission_mxn=Decimal(commission),
            ),
        )

    total = NetpayCommerceAttentionService._p2_totals([row("100", "10"), row("900", "45"), row("0", "20")])
    assert total.distributor_commission_mxn == 75 and total.effective_commission_on_sales_percent == Decimal("5.50")
    assert total.compatible_stores == 2 and total.monetary_stores == 3
    mixed = NetpayCommerceAttentionService._p2_totals([row("100", "10"), row("900", "45", "2026-10")])
    assert mixed.percentage_reason == "period_mismatch"


def test_synthetic_291_rows_repeated_headers_and_signed_control():
    """Synthetic structural fixture, deliberately not evidence about a real file."""
    from openpyxl import load_workbook

    w = load_workbook(BytesIO(workbook()))
    s = w.active
    s.delete_rows(8, s.max_row)
    h = list(next(s.iter_rows(min_row=7, max_row=7, values_only=True)))
    number = 0
    for entity, count in (("Asesor", 140), ("Asociado 15", 4), ("Distribuidor", 147)):
        if number:
            s.append(["NetPay"])
            s.append(["Comisiones Synthetic"])
            s.append(["Periodo Septiembre 2026"])
            s.append([])
            s.append(h)
        for _ in range(count):
            amount = "1" if number < 288 else "25380.64" if number == 288 else "-8187.85" if number == 289 else "0"
            s.append(
                [entity, "Synthetic", "Septiembre", None, "Synthetic", "UNKNOWN", amount, "synthetic@example.invalid"]
            )
            number += 1
    s.cell(2, 5, "17480.79")
    b = BytesIO()
    w.save(b)
    rows, _, meta = parse_settlement(b.getvalue())
    assert len(rows) == 291 and {r["source_header_row"] for r in rows} == {7, 152, 161}
    assert meta["positive"] == "25668.64" and meta["negative"] == "-8187.85" and meta["net_reported"] == "17480.79"
    assert meta["control_passed"] and rows[-1]["amount_signed"] == "0"
    s.cell(152, 6, "CONCEPTO ALTERADO")
    b = BytesIO()
    w.save(b)
    from fastapi import HTTPException

    with pytest.raises(HTTPException, match="settlement_row_invalid"):
        parse_settlement(b.getvalue())


def test_settlement_provenance_tamper_and_read_scope_boundaries():
    setup_data()
    b = settlement().json()
    assert accept(b).status_code == 200
    with app.state.yarvis.persistence.create_session() as db:
        batch = db.get(OperationalDataBatch, UUID(b["id"]))
        member = db.scalar(
            select(PrincipalMembership).where(PrincipalMembership.organization_id == batch.organization_id)
        )
        member.role = "netpay_inbox_portfolio_reader"
        db.commit()
    summary = client.get("/netpay/commerce-attention/settlements", headers=headers())
    assert summary.status_code == 200 and "movements" not in summary.json()[0]
    assert (
        client.get("/netpay/commerce-attention/settlements?include_details=true", headers=headers()).status_code == 403
    )
    with app.state.yarvis.persistence.create_session() as db:
        row = db.scalar(select(OperationalDataRow).where(OperationalDataRow.batch_id == UUID(b["id"])))
        row.controlled_payload = {**row.controlled_payload, "amount_signed": "999"}
        db.commit()
    assert client.get(f"/netpay/data/datasets/{b['id']}/results", headers=headers()).status_code == 409


def test_core_xlsx_raw_decimal_tokens_and_source_row():
    import zipfile

    w = Workbook()
    s = w.active
    s.title = "BD Rentabilidad"
    s.append(["Store ID", "Mes", "Volumen", "Rentabilidad", *COMPONENTS.values()])
    s.append([])
    s.append(["SYN-STORE-001", "2026-09", 1000, 700, 1, 0, 0])
    b = BytesIO()
    w.save(b)
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(b.getvalue())) as source, zipfile.ZipFile(out, "w") as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(
                    b'<c r="E3" t="n"><v>1</v>', b'<c r="E3" t="n"><v>0.123456789012345678901234</v>'
                )
            target.writestr(entry, content)
    store = setup_data()
    response = upload("precision.xlsx", out.getvalue(), "monthly_store_profitability", "precision")
    assert response.status_code == 201, response.text
    assert response.json()["rows"][0]["source_row_number"] == 3
    assert accept(response.json()).status_code == 200
    assert metric(store).distributor_commission_mxn == Decimal("0.123456789012345678901234")
