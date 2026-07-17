from io import BytesIO
from hashlib import sha256
from datetime import datetime

from fastapi.testclient import TestClient
from openpyxl import Workbook

from yarvis_api.main import app


client = TestClient(app)


def build_weekly_sales_xlsx() -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Weekly Summary"
    summary.append(["Weekly sales presentation"])

    detail = workbook.create_sheet("DETALLE")
    detail.append([
        "FECHA_CORTE", "store_id", "branch_name", "client_id", "TPV_HISTORICO",
        "ULTIMA_TX_DATE", "DIAS_SIN_TX", "TPV_MTD_CORTE", "TRX_MTD_CORTE", "UNMAPPED_COLUMN",
    ])
    detail.append([datetime(2026, 1, 1), "S-100", "Branch A", "C-1", 5000, datetime(2025, 11, 1), 61, 900, 10, "x"])
    detail.append([datetime(2026, 1, 1), "S-200", "Branch B", "C-2", 800, datetime(2025, 12, 30), 2, 200, 3, "y"])
    detail.append([None] * 10)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def upload(content: bytes, filename: str = "weekly_sales.xlsx"):
    return client.post(
        "/data-intake/documents",
        data={"source_type": "manual_upload", "source_name": "synthetic-test", "classification": "confidential"},
        files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def test_weekly_sales_xlsx_preview_is_deterministic_and_does_not_expose_rows():
    content = build_weekly_sales_xlsx()
    uploaded = upload(content)
    assert uploaded.status_code == 201, uploaded.text
    document_id = uploaded.json()["document_id"]

    document = client.get(f"/data-intake/documents/{document_id}")
    assert document.status_code == 200
    assert document.json()["file_hash"] == sha256(content).hexdigest()

    processed = client.post(f"/data-intake/documents/{document_id}/process")
    assert processed.status_code == 200, processed.text
    assert processed.json()["status"] == "preview_ready", processed.json()
    preview = client.get(f"/data-intake/documents/{document_id}/preview")
    assert preview.status_code == 200
    payload = preview.json()
    assert payload["detected_report_type"] == "netpay_weekly_sales_report"
    assert payload["worksheets"][0] == "DETALLE"
    assert payload["total_rows"] == 2
    assert payload["sample_rows"] == []
    assert payload["operational_summary"]["stores"] == 2
    assert payload["operational_summary"]["churn_candidates"] == 1
    assert any(item["source_header"] == "UNMAPPED_COLUMN" and item["canonical_field"] is None for item in payload["proposed_canonical_mappings"])


def test_confirmation_appends_confirmed_observations_and_updates_mission_control():
    uploaded = upload(build_weekly_sales_xlsx(), "another_weekly_sales.xlsx")
    document_id = uploaded.json()["document_id"]
    assert client.post(f"/data-intake/documents/{document_id}/process").status_code == 200

    before = client.get("/observations", params={"domain": "netpay", "confirmation_status": "candidate"})
    assert before.status_code == 200
    assert any(item["document_id"] == document_id for item in before.json())

    confirmed = client.post(f"/data-intake/documents/{document_id}/confirm", json={"reviewer": "synthetic-reviewer"})
    assert confirmed.status_code == 200
    assert confirmed.json()["review_status"] == "confirmed"

    after = client.get("/observations", params={"domain": "netpay", "confirmation_status": "confirmed"})
    assert any(item["document_id"] == document_id for item in after.json())
    mission = client.get("/mission-control/summary")
    assert mission.status_code == 200
    assert mission.json()["churn_candidates"] >= 1


def test_only_xlsx_is_accepted_and_duplicates_are_detected():
    content = build_weekly_sales_xlsx()
    first = upload(content, "first.xlsx")
    second = upload(content, "second.xlsx")
    assert first.status_code == 201
    assert second.status_code == 201
    assert second.json()["duplicate"] is True

    rejected = client.post(
        "/data-intake/documents",
        files={"file": ("not-allowed.csv", b"store_id\nS-1\n", "text/csv")},
    )
    assert rejected.status_code == 422


def test_churn_requires_supported_sixty_day_inactivity_and_never_uses_weekly_sales_alone():
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["presentation"])
    detail = workbook.create_sheet("DETALLE")
    detail.append(["FECHA_CORTE", "store_id", "branch_name", "client_id", "TPV_HISTORICO", "ULTIMA_TX_DATE", "DIAS_SIN_TX", "TPV_MTD_CORTE"])
    cutoff = datetime(2026, 1, 1)
    detail.append([cutoff, "S-ZERO", "A", "C-1", 1000, cutoff, 0, 0])
    detail.append([cutoff, "S-60", "A", "C-1", 1000, datetime(2025, 11, 2), 60, 100])
    detail.append([cutoff, "S-59", "A", "C-1", 1000, datetime(2025, 11, 3), 59, 100])
    detail.append([cutoff, "S-MISSING", "A", "C-1", 1000, None, 60, 100])
    detail.append([cutoff, "S-MALFORMED", "A", "C-1", 1000, "not-a-date", 60, 100])
    detail.append([cutoff, "S-STALE", "A", "C-1", 1000, datetime(2025, 12, 22), 60, 100])
    stream = BytesIO()
    workbook.save(stream)

    uploaded = upload(stream.getvalue(), "validation.xlsx")
    document_id = uploaded.json()["document_id"]
    assert client.post(f"/data-intake/documents/{document_id}/process").json()["status"] == "preview_ready"
    preview = client.get(f"/data-intake/documents/{document_id}/preview").json()
    assert preview["operational_summary"]["churn_candidates"] == 1
    assert preview["operational_summary"]["churn_insufficient_data"] == 3

    observations = client.get("/observations", params={"domain": "netpay", "confirmation_status": "candidate"}).json()
    findings = [item for item in observations if item["document_id"] == document_id and item["field_name"].startswith("finding.churn_")]
    assert [item["subject_reference"] for item in findings if item["field_name"] == "finding.churn_candidate"] == ["S-60"]
    assert {item["subject_reference"] for item in findings if item["field_name"] == "finding.churn_insufficient_data"} == {"S-MISSING", "S-MALFORMED", "S-STALE"}


def test_inactive_stores_report_derives_reviewable_findings_and_links_only_confirmed_store_id():
    weekly = upload(build_weekly_sales_xlsx(), "confirmed-weekly.xlsx")
    weekly_document_id = weekly.json()["document_id"]
    assert client.post(f"/data-intake/documents/{weekly_document_id}/process").status_code == 200
    assert client.post(f"/data-intake/documents/{weekly_document_id}/confirm", json={"reviewer": "synthetic-reviewer"}).status_code == 200

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inactive Stores"
    sheet.append(["#", "Store ID", "Store Name", "Master Name", "Distributor Name", "Estatus", "Meses No Uso", "Alertas"])
    sheet.append([1, "S-100", "Store A", "Master A", "Distributor A", "Active", 2, "Cobrar Renta o Cancelar"])
    sheet.append([2, "S-WATCH", "Store B", "Master B", "Distributor B", "Active", 1, ""])
    sheet.append([3, "S-BLOCK-A", "Store C", "Master C", "Distributor C", "Bloqueado terminal incorrecta", 0, ""])
    sheet.append([4, "S-BLOCK-B", "Store D", "Master D", "Distributor D", "Bloqueado cambio terminal", 0, ""])
    sheet.append([5, "S-ACTIVE", "Store E", "Master E", "Distributor E", "Active", 0, "Sin Activacion"])
    stream = BytesIO()
    workbook.save(stream)

    uploaded = upload(stream.getvalue(), "inactive-stores.xlsx")
    document_id = uploaded.json()["document_id"]
    assert client.post(f"/data-intake/documents/{document_id}/process").json()["status"] == "preview_ready"
    preview = client.get(f"/data-intake/documents/{document_id}/preview").json()
    assert preview["detected_report_type"] == "netpay_inactive_stores_report"
    assert preview["operational_summary"]["churn_watch"] == 1
    assert preview["operational_summary"]["churn_candidates"] == 1
    assert preview["operational_summary"]["cancellation_reviews"] == 1
    assert preview["operational_summary"]["operational_blocks"] == 2
    assert preview["operational_summary"]["activation_failures"] == 1
    assert preview["operational_summary"]["linked_weekly_sales"] == 1
    assert preview["operational_summary"]["critical_stores"] == 0
    assert {item["canonical_field"] for item in preview["proposed_canonical_mappings"] if item["canonical_field"]} == {"store_id", "store_name", "master_name", "distributor_name", "operational_status", "months_no_use", "source_alert"}

    findings = client.get("/observations", params={"domain": "netpay", "confirmation_status": "candidate"}).json()
    findings = [item for item in findings if item["document_id"] == document_id and item["field_name"].startswith("finding.")]
    assert {item["field_name"] for item in findings} >= {"finding.churn_watch", "finding.churn_candidate", "finding.cancellation_review", "finding.operational_block", "finding.activation_failure"}
    cancellation = next(item for item in findings if item["field_name"] == "finding.cancellation_review")
    assert cancellation["observed_value"]["recommended_action"] == "cancellation_review"
    assert cancellation["observed_value"]["requires_human_approval"] is True
    assert not any(item["field_name"] in {"company", "branch", "asset_serial", "finding.cancellation_executed"} for item in findings)
