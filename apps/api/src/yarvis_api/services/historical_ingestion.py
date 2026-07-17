"""Deterministic, unregistered ingestion for NetPay historical workbooks.

This module deliberately returns typed value objects only.  It neither persists
evidence nor logs source rows, so the route-registration gate remains closed.
"""
from __future__ import annotations

import hashlib
import math
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from yarvis_api.schemas.historical_evidence import (
    HistoricalEvidenceRecord,
    HistoricalHeaderResult,
    HistoricalSheetClassificationResult,
    HistoricalSheetInspection,
    HistoricalValidationError,
    HistoricalWorkbookInspection,
    ProfitabilityDuplicateGroup,
    ProfitabilityDuplicateResult,
    ProfitabilitySummary,
    TransactionHistoricalAggregate,
)

TRANSACTION_DETAIL_REPORT = "netpay_transaction_detail_report"
PROFITABILITY_REPORT = "netpay_profitability_report"
SUPPORTED_SHEETS = {
    "BD Detalle por Transaccion": TRANSACTION_DETAIL_REPORT,
    "BD Rentabilidad": PROFITABILITY_REPORT,
}

# Source spellings are intentionally retained here; matching is accent/case/space-insensitive.
TRANSACTION_DETAIL_MAPPING = {
    "Store ID": "store_id", "Store Name": "store_name_observed", "Client ID": "client_id",
    "Client Name": "client_name_observed", "Razón Social": "legal_company_name_observed",
    "UEN": "business_unit", "ID Distribuidor": "distributor_id", "Producto": "product_observed",
    "Giro Natural": "merchant_category_observed", "Vales": "vouchers_supported",
    "Tipo de Transacción": "transaction_type", "Banco Adquirente": "acquiring_bank",
    "Banco Emisor": "issuing_bank", "Procesador": "processor", "Mes": "reporting_month",
    "Tarjeta": "card_type", "MSI": "installment_months", "Tasa de Venta": "sale_rate",
    "Tasa de Compra": "purchase_rate", "Transacciones": "transaction_count",
    "Volumen": "transaction_volume", "Ingreso": "income", "Costo": "cost",
    "Diferencial de Tasas": "rate_margin",
}
PROFITABILITY_MAPPING = {
    "Store ID": "store_id", "Store Name": "store_name_observed", "Client ID": "client_id",
    "Client Name": "client_name_observed", "Branch": "branch_name_observed", "UEN": "business_unit",
    "Mes Creación Store": "store_creation_month", "ID Distribuidor": "distributor_id",
    "Distribuidor": "distributor_name_observed", "ID Asesor": "advisor_id", "Asesor": "advisor_name_observed",
    "ID Asociado": "associate_id", "Asociado": "associate_name_observed", "ID Referenciador": "referrer_id",
    "Mes": "reporting_month", "Store IDs sin Uso": "no_use_indicator_historical",
    "Volumen": "sales_volume", "Volumen Tiempo Aire": "airtime_volume", "Diferencial de Tasas": "rate_margin",
    "Renta Pagada": "rent_paid", "Total Comisiones (sin Activación ni Referenciador)": "total_commissions",
    "Rentabilidad": "profitability", "Master": "master_name_observed",
}
REQUIRED_HEADERS = {
    TRANSACTION_DETAIL_REPORT: {"Store ID", "Client ID", "Tipo de Transacción", "Transacciones", "Volumen", "Tasa de Venta", "Tasa de Compra", "Diferencial de Tasas"},
    PROFITABILITY_REPORT: {"Store ID", "Client ID", "Branch", "Volumen", "Renta Pagada", "Total Comisiones (sin Activación ni Referenciador)", "Rentabilidad"},
}
IDENTIFIER_FIELDS = {"store_id", "client_id", "distributor_id", "advisor_id", "associate_id", "referrer_id"}
NUMERIC_FIELDS = {
    "transaction_count", "transaction_volume", "income", "cost", "rate_margin", "sale_rate", "purchase_rate",
    "sales_volume", "airtime_volume", "rent_paid", "total_commissions", "profitability",
}


def _error(code: str, message: str, field: str | None = None) -> HistoricalValidationError:
    return HistoricalValidationError("", None, field, code, "error", message)


def _at_source(error: HistoricalValidationError, sheet: str, row: str | None, field: str | None = None) -> HistoricalValidationError:
    return HistoricalValidationError(sheet, row, field if field is not None else error.source_field, error.error_code, error.severity, error.message)


def normalize_header(value: Any) -> str:
    """Normalize labels, never source values, for deterministic header matching."""
    text = unicodedata.normalize("NFD", str(value or "").strip()).casefold()
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"\s+", " ", text)


def _mapping_for(report: str) -> dict[str, str]:
    source_mapping = TRANSACTION_DETAIL_MAPPING if report == TRANSACTION_DETAIL_REPORT else PROFITABILITY_MAPPING
    return {normalize_header(source): canonical for source, canonical in source_mapping.items()}


def classify_historical_sheet(headers: list[Any]) -> HistoricalSheetClassificationResult:
    normalized = {normalize_header(header) for header in headers if normalize_header(header)}
    matches = [
        report for report, required in REQUIRED_HEADERS.items()
        if {normalize_header(header) for header in required}.issubset(normalized)
    ]
    conflicts: list[str] = []
    if not matches:
        conflicts.append("missing_required_header")
    elif len(matches) > 1:
        conflicts.append("ambiguous_sheet_classification")
    return HistoricalSheetClassificationResult(matches[0] if len(matches) == 1 else None, matches, conflicts, len(matches) == 1)


def normalize_identifier(value: Any) -> tuple[str | None, list[HistoricalValidationError]]:
    if value is None or isinstance(value, bool):
        return None, []
    if isinstance(value, str):
        result = value.strip()
        if not result or result.casefold() in {"none", "nan", "null"}:
            return None, []
        # Do not convert textual identifiers: leading zeros and intentional punctuation matter.
        return result, []
    if isinstance(value, Decimal):
        if not value.is_finite() or value != value.to_integral_value():
            return None, [_error("invalid_store_id", "Identifier must be an integer-like value.")]
        return format(value.quantize(Decimal("1")), "f"), []
    if isinstance(value, int):
        return str(value), []
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            return None, [_error("invalid_store_id", "Identifier must be an integer-like value.")]
        return str(int(value)), []
    return None, [_error("invalid_store_id", "Identifier has an unsupported type.")]


def normalize_reporting_month(value: Any) -> tuple[str | None, list[HistoricalValidationError]]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, []
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m"), []
    if isinstance(value, str):
        text = value.strip()
        match = re.fullmatch(r"(\d{4})-(0[1-9]|1[0-2])", text)
        if match:
            return text, []
        match = re.fullmatch(r"(\d{4})-(0[1-9]|1[0-2])-[0-3]\d", text)
        if match:
            return f"{match.group(1)}-{match.group(2)}", []
    return None, [_error("invalid_reporting_period", "Reporting period must be a complete supported month representation.")]


def normalize_decimal(value: Any) -> tuple[Decimal | None, list[HistoricalValidationError]]:
    if value is None or isinstance(value, bool):
        return None, []
    if isinstance(value, Decimal):
        return (value, []) if value.is_finite() else (None, [_error("invalid_numeric_value", "Numeric value must be finite.")])
    if isinstance(value, int):
        return Decimal(value), []
    if isinstance(value, float):
        return (Decimal(str(value)), []) if math.isfinite(value) else (None, [_error("invalid_numeric_value", "Numeric value must be finite.")])
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None, []
        if "%" in text:
            return None, [_error("invalid_numeric_value", "Percentage notation is ambiguous for this numeric field.")]
        normalized = text.replace("$", "").replace(",", "").replace(" ", "")
        try:
            candidate = Decimal(normalized)
        except InvalidOperation:
            return None, [_error("invalid_numeric_value", "Numeric value is not in a supported format.")]
        if not candidate.is_finite():
            return None, [_error("invalid_numeric_value", "Numeric value must be finite.")]
        return candidate, []
    return None, [_error("invalid_numeric_value", "Numeric value has an unsupported type.")]


def _normalize_rate(value: Any) -> tuple[Decimal | None, list[HistoricalValidationError]]:
    """Accept explicit percentages only for fields whose headers establish rate semantics."""
    if isinstance(value, str) and "%" in value:
        return normalize_decimal(value.replace("%", ""))
    return normalize_decimal(value)


def _header_results(headers: list[Any], report: str | None) -> tuple[list[HistoricalHeaderResult], list[str]]:
    normalized = [normalize_header(header) for header in headers]
    collisions = sorted({key for key in normalized if key and normalized.count(key) > 1})
    mapping = _mapping_for(report) if report else {}
    results = [HistoricalHeaderResult(str(header or ""), key, mapping.get(key), key in mapping, key in collisions) for header, key in zip(headers, normalized)]
    return results, collisions


def _sheet_periods(sheet: Any, headers: list[Any], promotable: bool) -> list[str]:
    if not promotable:
        return []
    period_index = next((i for i, header in enumerate(headers) if normalize_header(header) == normalize_header("Mes")), None)
    if period_index is None:
        return []
    periods = set()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        period, _ = normalize_reporting_month(row[period_index] if len(row) > period_index else None)
        if period:
            periods.add(period)
    return sorted(periods)


def inspect_historical_workbook(path: str | Path) -> HistoricalWorkbookInspection:
    workbook_path = Path(path)
    fingerprint = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    results: list[HistoricalSheetInspection] = []
    supported: list[str] = []
    ignored: list[str] = []
    unresolved: list[str] = []
    for name in workbook.sheetnames:
        sheet = workbook[name]
        headers = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
        expected = SUPPORTED_SHEETS.get(name)
        classification_result = classify_historical_sheet(headers) if expected else HistoricalSheetClassificationResult(None, [], [], False)
        headers_result, collisions = _header_results(headers, classification_result.classification)
        structural_conflicts = list(classification_result.structural_conflicts)
        if collisions:
            structural_conflicts.append("normalized_header_collision")
        classification = classification_result.classification if expected == classification_result.classification and not collisions else None
        promotable = classification is not None
        if expected is None:
            ignored.append(name)
        elif promotable:
            supported.append(name)
        else:
            unresolved.append(name)
        candidate_rows = sum(1 for row in sheet.iter_rows(min_row=2, values_only=True) if any(value is not None and str(value).strip() for value in row))
        results.append(HistoricalSheetInspection(
            source_sheet=name, classification=classification, source_column_count=len(headers), candidate_row_count=candidate_rows,
            mapped_headers=[item for item in headers_result if item.mapped], unresolved_headers=[item for item in headers_result if not item.mapped],
            normalized_header_collisions=collisions, structural_conflicts=structural_conflicts,
            reporting_periods=_sheet_periods(sheet, headers, promotable), promotable=promotable,
        ))
    workbook.close()
    return HistoricalWorkbookInspection(fingerprint, workbook_path.name, list(workbook.sheetnames), supported, ignored, unresolved, results, bool(supported) and not unresolved)


def normalize_historical_workbook(path: str | Path, ingestion_timestamp: datetime) -> tuple[HistoricalWorkbookInspection, list[HistoricalEvidenceRecord]]:
    inspection = inspect_historical_workbook(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    evidence: list[HistoricalEvidenceRecord] = []
    for result in inspection.sheet_results:
        if not result.promotable or not result.classification:
            continue
        sheet = workbook[result.source_sheet]
        headers = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
        mapping = _mapping_for(result.classification)
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            canonical: dict[str, Any] = {}
            unresolved: dict[str, Any] = {}
            original_headers: dict[str, str] = {}
            errors: list[HistoricalValidationError] = []
            row_reference = f"row:{row_number}"
            for header, value in zip(headers, row):
                source_header = str(header or "")
                field = mapping.get(normalize_header(header))
                if field is None:
                    unresolved[source_header] = value
                    continue
                original_headers[field] = source_header
                if field in IDENTIFIER_FIELDS:
                    normalized, field_errors = normalize_identifier(value)
                    if field_errors and field != "store_id":
                        field_errors = [_error("invalid_client_id", "Client identifier must be an integer-like value.", field)]
                elif field in {"reporting_month", "store_creation_month"}:
                    normalized, field_errors = normalize_reporting_month(value)
                elif field in {"sale_rate", "purchase_rate", "rate_margin"}:
                    normalized, field_errors = _normalize_rate(value)
                elif field in NUMERIC_FIELDS:
                    normalized, field_errors = normalize_decimal(value)
                else:
                    normalized, field_errors = value, []
                canonical[field] = normalized
                errors.extend(_at_source(error, result.source_sheet, row_reference, source_header) for error in field_errors)
            if not canonical.get("store_id"):
                errors.append(HistoricalValidationError(result.source_sheet, row_reference, original_headers.get("store_id"), "missing_store_id", "error", "Store identifier is required for store-level consolidation."))
            if not canonical.get("reporting_month"):
                errors.append(HistoricalValidationError(result.source_sheet, row_reference, original_headers.get("reporting_month"), "missing_reporting_period", "error", "Reporting period is required for historical consolidation."))
            evidence.append(HistoricalEvidenceRecord(
                evidence_type=result.classification, temporal_class="historical", canonical_fields=canonical,
                unresolved_source_fields=unresolved, document_fingerprint=inspection.document_fingerprint,
                source_workbook=inspection.source_workbook, source_sheet=result.source_sheet,
                source_row_reference=row_reference, original_headers=original_headers,
                reporting_period=canonical.get("reporting_month"), ingestion_timestamp=ingestion_timestamp,
                validation_errors=errors, identity_review_required=not bool(canonical.get("store_id")),
            ))
    workbook.close()
    return inspection, evidence


def _provenance(record: HistoricalEvidenceRecord) -> str:
    return f"{record.source_sheet}:{record.source_row_reference}"


def _sum_decimal(records: list[HistoricalEvidenceRecord], field: str) -> Decimal | None:
    values = [record.canonical_fields.get(field) for record in records if isinstance(record.canonical_fields.get(field), Decimal)]
    return sum(values, Decimal("0")) if values else None


def _observed(records: list[HistoricalEvidenceRecord], field: str) -> list[str]:
    return sorted({value.strip() for record in records if isinstance((value := record.canonical_fields.get(field)), str) and value.strip()})


def aggregate_transaction_evidence(evidence_records: list[HistoricalEvidenceRecord]) -> tuple[list[TransactionHistoricalAggregate], list[HistoricalEvidenceRecord]]:
    groups: dict[tuple[str, str], list[HistoricalEvidenceRecord]] = defaultdict(list)
    excluded: list[HistoricalEvidenceRecord] = []
    for record in evidence_records:
        store_id = record.canonical_fields.get("store_id")
        if record.evidence_type != TRANSACTION_DETAIL_REPORT or not isinstance(store_id, str) or not record.reporting_period:
            excluded.append(record)
            continue
        groups[(store_id, record.reporting_period)].append(record)
    aggregates: list[TransactionHistoricalAggregate] = []
    for (store_id, period), records in sorted(groups.items()):
        # The source header is "Diferencial de Tasas".  It is rate-like/ambiguous rather
        # than proven monetary, so preserve observations and never sum it.
        rate_values = sorted({value for record in records if isinstance((value := record.canonical_fields.get("rate_margin")), Decimal)})
        aggregates.append(TransactionHistoricalAggregate(
            store_id, period, _sum_decimal(records, "transaction_count"), _sum_decimal(records, "transaction_volume"),
            _sum_decimal(records, "income"), _sum_decimal(records, "cost"), None, rate_values,
            _observed(records, "product_observed"), _observed(records, "transaction_type"), _observed(records, "card_type"),
            _observed(records, "acquiring_bank"), _observed(records, "issuing_bank"), _observed(records, "processor"),
            len(records), sorted(_provenance(record) for record in records),
            [error for record in records for error in record.validation_errors],
        ))
    return aggregates, excluded


def detect_profitability_duplicates(evidence_records: list[HistoricalEvidenceRecord]) -> ProfitabilityDuplicateResult:
    groups: dict[tuple[str, str], list[HistoricalEvidenceRecord]] = defaultdict(list)
    excluded: list[HistoricalEvidenceRecord] = []
    for record in evidence_records:
        store_id = record.canonical_fields.get("store_id")
        if record.evidence_type != PROFITABILITY_REPORT or not isinstance(store_id, str) or not record.reporting_period:
            if record.evidence_type == PROFITABILITY_REPORT:
                excluded.append(record)
            continue
        groups[(store_id, record.reporting_period)].append(record)
    accepted: list[ProfitabilitySummary] = []
    exact: list[ProfitabilityDuplicateGroup] = []
    conflicting: list[ProfitabilityDuplicateGroup] = []
    for (store_id, period), records in sorted(groups.items()):
        provenance = sorted(_provenance(record) for record in records)
        if len(records) == 1:
            accepted.append(ProfitabilitySummary(store_id, period, records[0].canonical_fields, provenance))
        elif all(record.canonical_fields == records[0].canonical_fields for record in records[1:]):
            accepted.append(ProfitabilitySummary(store_id, period, records[0].canonical_fields, provenance))
            exact.append(ProfitabilityDuplicateGroup(store_id, period, "exact", len(records), provenance))
        else:
            conflicting.append(ProfitabilityDuplicateGroup(store_id, period, "conflict", len(records), provenance))
    return ProfitabilityDuplicateResult(accepted, exact, conflicting, excluded)
