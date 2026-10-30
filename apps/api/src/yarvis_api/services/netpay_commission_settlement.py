"""IG-015 source parsing and settlement lifecycle, owned by Netpay.

No concept mapping is certified by this profile. Original amounts and entities
remain evidence; acceptance is not bank confirmation or Core reconciliation.
"""

import hashlib
import json
import re
import unicodedata
import zipfile
from datetime import date
from decimal import Decimal, InvalidOperation, localcontext
from io import BytesIO
from xml.etree import ElementTree as ET

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from yarvis_api.models.netpay_master import NetpayStoreReference
from yarvis_api.models.netpay_operational_data import OperationalDataBatch, OperationalDataRow

MAX_PACKAGE_ENTRIES = 500
MAX_PACKAGE_ENTRY_BYTES = 20_000_000
MAX_PACKAGE_EXPANDED_BYTES = 50_000_000
MAX_PACKAGE_COMPRESSION_RATIO = 200
ALLOWED_WORKBOOK_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
COMPONENTS = {
    "activation_bonus": "Bono de Activación",
    "referrer_commission": "Comisión - Referenciador",
    "base_commission": "Total Comisiones (sin Activación ni Referenciador)",
}
DEFINITION = "ap_aq_ar_v2"
MONTHS = {
    name: i
    for i, name in enumerate(
        "enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre".split(), 1
    )
}
HEADERS = ["Entidad", "NOMBRE", "Mes", "Store", "Comercio", "CONCEPTO", "Total Comisiones", "CORREO"]


def normal(value):
    return re.sub(
        r"[^a-z0-9]", "", unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    )


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"), default=str).encode()).hexdigest()


def decimal_value(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        result = Decimal(str(value))
        if not result.is_finite() or len(result.as_tuple().digits) > 38 or abs(result.as_tuple().exponent) > 24:
            raise ValueError()
        return result
    except (InvalidOperation, ValueError):
        raise HTTPException(422, "invalid_decimal_source") from None


def exact_sum(values):
    # Bounded source values (38 digits / 24 scale) and <=5000 rows fit exactly.
    with localcontext() as context:
        context.prec = 90
        return sum(values, Decimal(0))


def concept_candidate(concept):
    """A lexical candidate is NOT an approved AP/AQ/AR correspondence."""
    label = normal(concept)
    if label == "comisionporactivacion":
        return "activation_candidate"
    if label.startswith("pagocomisiones"):
        return "base_commission_candidate"
    if label.startswith(
        (
            "revcomerciosnorentables",
            "comerciosnorentables",
            "penalizacion",
            "servicioasesorinterno",
            "cobroasesorinterno",
        )
    ):
        return "adjustment_candidate"
    return "unclassified"


def _invalid_workbook(detail="invalid_settlement_workbook"):
    raise HTTPException(422, detail)


def _canonical_package_name(name):
    if not name or "\x00" in name or name.startswith(("/", "\\")) or re.match(r"^[A-Za-z]:", name):
        _invalid_workbook("unsafe_workbook_path")
    canonical = name.replace("\\", "/")
    parts = canonical.split("/")
    if any(part in ("", ".", "..") for part in parts):
        _invalid_workbook("unsafe_workbook_path")
    return canonical


def normalize_ooxml_package(content):
    """Validate OOXML and canonicalize only ZIP entry separators in memory."""
    try:
        source = zipfile.ZipFile(BytesIO(content))
    except (zipfile.BadZipFile, ValueError):
        _invalid_workbook()
    with source:
        entries = source.infolist()
        if not entries or len(entries) > MAX_PACKAGE_ENTRIES:
            _invalid_workbook("workbook_package_limits_exceeded")
        originals, canonical_names, casefolded = set(), set(), set()
        expanded = 0
        normalized = []
        for entry in entries:
            if entry.filename in originals:
                _invalid_workbook("duplicate_workbook_entry")
            originals.add(entry.filename)
            if entry.is_dir() or entry.flag_bits & 0x1:
                _invalid_workbook("unsafe_workbook_entry")
            if entry.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                _invalid_workbook("unsafe_workbook_compression")
            if entry.file_size > MAX_PACKAGE_ENTRY_BYTES:
                _invalid_workbook("workbook_package_limits_exceeded")
            expanded += entry.file_size
            if expanded > MAX_PACKAGE_EXPANDED_BYTES:
                _invalid_workbook("workbook_package_limits_exceeded")
            if entry.file_size and (
                not entry.compress_size or entry.file_size / entry.compress_size > MAX_PACKAGE_COMPRESSION_RATIO
            ):
                _invalid_workbook("workbook_compression_ratio_exceeded")
            canonical = _canonical_package_name(entry.filename)
            folded = canonical.casefold()
            if canonical in canonical_names or folded in casefolded:
                _invalid_workbook("workbook_entry_collision")
            canonical_names.add(canonical)
            casefolded.add(folded)
            normalized.append((entry, canonical))

        if any(name.startswith("xl/externalLinks/") or name.endswith("vbaProject.bin") for name in canonical_names):
            _invalid_workbook("unsafe_workbook")
        required = {"[Content_Types].xml", "xl/workbook.xml"}
        if not required.issubset(canonical_names):
            _invalid_workbook()

        payloads = {}
        try:
            for entry, canonical in normalized:
                payloads[canonical] = source.read(entry)
        except (RuntimeError, zipfile.BadZipFile, OSError):
            _invalid_workbook()

        try:
            types = ET.fromstring(payloads["[Content_Types].xml"])
            workbook_types = [
                node.attrib.get("ContentType") for node in types if node.attrib.get("PartName") == "/xl/workbook.xml"
            ]
            if workbook_types != [ALLOWED_WORKBOOK_CONTENT_TYPE]:
                _invalid_workbook("unsupported_workbook_type")
            for name, payload in payloads.items():
                if not name.endswith(".rels"):
                    continue
                relationships = ET.fromstring(payload)
                if any(
                    node.attrib.get("TargetMode", "").lower() == "external"
                    or re.match(r"^[A-Za-z][A-Za-z0-9+.-]*:", node.attrib.get("Target", ""))
                    or node.attrib.get("Target", "").startswith("//")
                    for node in relationships
                ):
                    _invalid_workbook("external_workbook_relationship")
        except ET.ParseError:
            _invalid_workbook()

        if all(entry.filename == canonical for entry, canonical in normalized):
            return content
        target_bytes = BytesIO()
        with zipfile.ZipFile(target_bytes, "w") as target:
            for entry, canonical in normalized:
                clone = zipfile.ZipInfo(canonical, entry.date_time)
                clone.comment = entry.comment
                clone.extra = entry.extra
                clone.internal_attr = entry.internal_attr
                clone.external_attr = entry.external_attr
                clone.create_system = entry.create_system
                clone.compress_type = entry.compress_type
                target.writestr(clone, payloads[canonical])
        return target_bytes.getvalue()


def exact_rows(content, title):
    """Keep date cells typed while replacing numeric floats with XML Decimal tokens."""
    content = normalize_ooxml_package(content)
    with zipfile.ZipFile(BytesIO(content)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 50_000_000:
            raise HTTPException(422, "expanded_workbook_limit")
        if any("externalLinks/" in n or n.endswith("vbaProject.bin") for n in archive.namelist()):
            raise HTTPException(422, "unsafe_workbook")
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        wb = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = {r.attrib["Id"]: r.attrib["Target"] for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        s = next(s for s in wb.find("s:sheets", ns) if s.attrib["name"] == title)
        target = rels[s.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]]
        target = target.lstrip("/") if target.startswith("/") else "xl/" + target
        tokens = {}
        for c in ET.fromstring(archive.read(target)).findall("s:sheetData/s:row/s:c", ns):
            if c.find("s:f", ns) is not None:
                raise HTTPException(422, "formulas_not_allowed")
            v = c.find("s:v", ns)
            if c.attrib.get("t") in (None, "n") and v is not None and v.text:
                tokens[c.attrib["r"]] = decimal_value(v.text)
    # Normal mode derives bounds when the producer omits the optional dimension.
    workbook = load_workbook(BytesIO(content), read_only=False, data_only=False, keep_links=False)
    try:
        result = []
        for row in workbook[title].iter_rows():
            if len(result) >= 5000 or len(row) > 50:
                raise HTTPException(422, "dataset_limits_exceeded")
            result.append(
                [
                    tokens.get(c.coordinate, c.value)
                    if hasattr(c, "coordinate") and not getattr(c, "is_date", False)
                    else c.value
                    for c in row
                ]
            )
        return result
    finally:
        workbook.close()


def parse_settlement(content):
    try:
        workbook_content = normalize_ooxml_package(content)
        workbook = load_workbook(BytesIO(workbook_content), read_only=False, keep_links=False)
        names = workbook.sheetnames
        workbook.close()
        if len(names) != 1:
            raise HTTPException(422, "settlement_requires_single_sheet")
        rows = exact_rows(workbook_content, names[0])
    except (zipfile.BadZipFile, KeyError, ValueError, ET.ParseError):
        raise HTTPException(422, "invalid_settlement_workbook") from None
    if len(rows) < 8:
        raise HTTPException(422, "settlement_headers_missing")
    header = [str(v or "") for v in rows[0]]
    if header != ["NOMBRE", "", "CORREO", "", "Total Comisiones", "Mes Corriente", "Mes Previo", "Mensaje"]:
        raise HTTPException(422, "settlement_summary_header_invalid")
    announced = decimal_value(rows[1][4])
    if announced is None:
        raise HTTPException(422, "settlement_total_missing")
    periods = set()
    for row in rows:
        match = re.fullmatch(r"Periodo\s+(\w+)\s+(20\d{2})", str(row[0] or ""), re.I)
        if match and normal(match[1]) in MONTHS:
            periods.add(f"{match[2]}-{MONTHS[normal(match[1])]:02d}")
    if len(periods) != 1:
        raise HTTPException(422, "settlement_period_conflict")
    period = periods.pop()
    message = str(rows[1][7])
    payment = re.search(r"pago del (\d{1,2}) de (\w+)(?: de (20\d{2}))?", message, re.I)
    if not payment or normal(payment[2]) not in MONTHS:
        raise HTTPException(422, "announced_payment_date_missing")
    # A relative year is accepted only when explicitly tied to the same year.
    if not payment[3] and ("año en curso" not in message.lower() or period[:4] not in message):
        raise HTTPException(422, "announced_payment_year_ambiguous")
    if not payment[3] and MONTHS[normal(payment[2])] < int(period[5:]):
        raise HTTPException(422, "announced_payment_year_ambiguous")
    try:
        payment_date = date(int(payment[3] or period[:4]), MONTHS[normal(payment[2])], int(payment[1])).isoformat()
    except ValueError:
        raise HTTPException(422, "announced_payment_date_invalid") from None
    movements, block = [], None
    for number, original in enumerate(rows[2:], 3):
        row = (original + [None] * 8)[:8]
        if [str(v or "") for v in row] == HEADERS:
            block = number
            continue
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        first = str(row[0] or "")
        if all(v in (None, "") for v in row[1:]) and (
            first == "NetPay" or first.startswith("Comisiones ") or first.startswith("Periodo ")
        ):
            continue
        if (
            not block
            or not row[0]
            or not row[5]
            or normal(row[2]) not in MONTHS
            or MONTHS[normal(row[2])] != int(period[5:])
        ):
            raise HTTPException(422, f"settlement_row_invalid:{number}")
        amount = decimal_value(row[6])
        if amount is None:
            raise HTTPException(422, f"settlement_amount_missing:{number}")
        store = None if row[3] in (None, "") else str(row[3]).strip()
        if store and not re.fullmatch(r"[A-Za-z0-9_-]{1,100}", store):
            raise HTTPException(422, f"settlement_store_invalid:{number}")
        movements.append(
            {
                "source_row_number": number,
                "source_header_row": block,
                "entity_original": str(row[0]),
                "concept_original": str(row[5]),
                "store_id": store,
                "reporting_period": period,
                "amount_signed": str(amount),
                "currency_code": "MXN",
            }
        )
    if not movements:
        raise HTTPException(422, "settlement_empty")
    amounts = [Decimal(r["amount_signed"]) for r in movements]
    net = exact_sum(amounts)
    metadata = {
        "announced_total": str(announced),
        "net_reported": str(net),
        "positive": str(exact_sum(x for x in amounts if x > 0)),
        "negative": str(exact_sum(x for x in amounts if x < 0)),
        "control_passed": net == announced,
        "announced_payment_date": payment_date,
        "bank_status": "not_confirmed",
        "currency_basis": "ratified_MXN_profile",
        "reporting_period": period,
        "mapping_version": "unconfirmed_v1",
    }
    return movements, names[0], metadata


def match_store(db, org, store):
    if not store:
        return None, "unmatched"
    refs = db.scalars(
        select(NetpayStoreReference).where(
            NetpayStoreReference.organization_id == org,
            NetpayStoreReference.store_id == store,
            NetpayStoreReference.active.is_(True),
        )
    ).all()
    return (refs[0].id, "matched") if len(refs) == 1 else (None, "ambiguous" if refs else "unmatched")


def stage(db, current, movements, sheet, metadata, source_hash):
    existing = db.scalar(
        select(OperationalDataBatch).where(
            OperationalDataBatch.organization_id == current.organization_id,
            OperationalDataBatch.dataset_type == "commission_settlement",
            OperationalDataBatch.source_hash == source_hash,
        )
    )
    if existing:
        return existing, 200
    batch = OperationalDataBatch(
        organization_id=current.organization_id,
        dataset_type="commission_settlement",
        reporting_period=metadata["reporting_period"],
        selected_sheet=sheet,
        sanitized_filename="commission-settlement.xlsx",
        source_hash=source_hash,
        source_profile="netpay_commission_settlement_v1",
        currency_code="MXN",
        status="needs_review",
        store_state_hash="0" * 64,
        preview_token=digest([source_hash, sheet, metadata, movements]),
        settlement_metadata=metadata,
        uploaded_by_principal_id=current.principal_id,
        row_counts={"received": len(movements), "authorized": len(movements), "invalid": 0},
    )
    db.add(batch)
    db.flush()
    for item in movements:
        reference, match = match_store(db, current.organization_id, item["store_id"])
        db.add(
            OperationalDataRow(
                batch_id=batch.id,
                source_row_number=item["source_row_number"],
                controlled_payload=item,
                validation_status="valid",
                match_status=match,
                store_reference_id=reference,
                error_codes=[] if reference else ["store_missing" if not item["store_id"] else "store_unmatched"],
                row_key=digest([sheet, item["source_row_number"]]),
                row_fingerprint=digest(item),
                projected_action="insert",
            )
        )
    return batch, 201


def verified_rows(db, batch):
    rows = db.scalars(
        select(OperationalDataRow)
        .where(OperationalDataRow.batch_id == batch.id)
        .order_by(OperationalDataRow.source_row_number)
    ).all()
    material = [r.controlled_payload for r in rows]
    if (
        any(digest(r.controlled_payload) != r.row_fingerprint for r in rows)
        or digest([batch.source_hash, batch.selected_sheet, batch.settlement_metadata, material]) != batch.preview_token
    ):
        raise HTTPException(409, "accepted_provenance_unverified")
    return rows


def accept(db, current, batch, payload, lock):
    rows = verified_rows(db, batch)
    if not batch.settlement_metadata.get("control_passed"):
        raise HTTPException(409, "settlement_control_difference")
    lock(db, ["settlement-effective", str(current.organization_id), batch.reporting_period])
    previous = db.scalar(
        select(OperationalDataBatch)
        .where(
            OperationalDataBatch.organization_id == current.organization_id,
            OperationalDataBatch.dataset_type == "commission_settlement",
            OperationalDataBatch.reporting_period == batch.reporting_period,
            OperationalDataBatch.settlement_effective.is_(True),
        )
        .with_for_update()
    )
    if previous:
        if payload.expected_predecessors.get("settlement") != previous.id or not payload.correction_reason:
            raise HTTPException(409, "explicit_predecessor_and_reason_required")
        previous.settlement_effective = False
        batch.predecessor_batch_id = previous.id
        db.flush()
    elif payload.expected_predecessors:
        raise HTTPException(409, "settlement_predecessor_conflict")
    for row in rows:
        row.store_reference_id, row.match_status = match_store(
            db, current.organization_id, row.controlled_payload["store_id"]
        )
        row.error_codes = (
            []
            if row.store_reference_id
            else ["store_missing" if not row.controlled_payload["store_id"] else "store_unmatched"]
        )
    batch.settlement_effective, batch.status = True, "accepted"
    return batch, 200
