from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
import os
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from io import StringIO
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from yarvis_api.clock import utc_now
from yarvis_api.models.domain_event import record_event
from yarvis_api.models.observation_engine import DocumentRecord, Observation, SourceRecord
from yarvis_api.modules.document_intelligence import DocumentIntelligenceProvider
from yarvis_api.storage.document_repository import LocalDocumentRepository

ALLOWED_EXTENSIONS = {".xlsx"}
REJECTED_EXTENSIONS = {".exe", ".bat", ".cmd", ".com", ".dll", ".msi", ".zip", ".rar", ".7z", ".tar", ".gz", ".xlsm"}

NETPAY_HEADER_MAPPING: dict[str, list[str]] = {
    "client_id": ["client id", "clientid", "numero de cliente", "número de cliente"],
    "company_name": ["company", "razon social", "razón social", "merchant", "comercio"],
    "rfc": ["rfc"],
    "branch_name": ["branch", "sucursal", "branch name"],
    "store_id": ["store id", "storeid", "numero de tienda", "número de tienda", "no. comercio", "no comercio"],
    "store_type": ["store type", "tipo de tienda"],
    "product_type": ["product", "product type", "tipo producto"],
    "asset_serial": ["serial", "device serial", "serie", "asset serial"],
    "terminal_model": ["model", "terminal model", "modelo"],
    "sales_volume": ["sales", "sales volume", "volumen de ventas", "monto"],
    "transaction_count": ["transactions", "transaction count", "numero transacciones", "número transacciones"],
    "last_activity_at": ["last activity", "ultima actividad", "última actividad"],
    "inactive_days": ["inactive days", "dias inactivos", "días inactivos"],
    "status": ["status", "estado"],
    "tracking_number": ["tracking", "guia", "guía", "tracking number"],
    "service_case_folio": ["folio", "service case", "caso"],
    "address": ["address", "domicilio"],
    "city": ["city", "ciudad"],
    "state": ["state", "estado"],
    "postal_code": ["postal code", "codigo postal", "código postal", "cp"],
    "contact_name": ["contact", "contact name", "contacto", "destinatario"],
    "phone": ["phone", "telefono", "teléfono"],
    "email": ["email", "correo"],
}


NETPAY_WEEKLY_SALES_MAPPING: dict[str, str] = {
    "FECHA_CORTE": "report_cutoff_at",
    "FECHA_REF_ESTATUS": "status_reference_at",
    "DISTRIBUIDOR": "distributor_name",
    "distributor_id": "distributor_id",
    "ASOCIADO": "associate_name",
    "ASESOR": "advisor_name",
    "store_id": "store_id",
    "status": "store_status",
    "UEN": "business_unit",
    "Fecha Creacion": "store_created_at",
    "Fecha Aprobacion": "store_approved_at",
    "Fecha Cancelacion": "store_cancelled_at",
    "Comercio": "commercial_name",
    "store_name": "store_name",
    "branch_name": "branch_name",
    "tax_reference": "tax_reference",
    "client_id": "client_id",
    "client_name": "client_name",
    "Giro": "business_category",
    "estado": "state",
    "ciudad": "city",
    "TPV_HISTORICO": "historical_sales_volume",
    "ULTIMA_TX_DATE": "last_activity_at",
    "DIAS_SIN_TX": "inactive_days",
    "TPV_MTD_CORTE": "sales_volume",
    "TRX_MTD_CORTE": "transaction_count",
    "FLAG_CONVERTIDO_ACTIVO": "active_conversion_flag",
    "FLAG_DORMIDO": "dormant_flag",
    "FLAG_PRECHURN": "prechurn_flag",
    "FLAG_CHURNEADO": "churned_flag",
    "ESTATUS_DEVICE_QS": "device_status",
    "ESTATUS_CALCULADO": "calculated_status",
}

NETPAY_INACTIVE_STORES_MAPPING: dict[str, str] = {
    "Store ID": "store_id",
    "Store Name": "store_name",
    "Master Name": "master_name",
    "Distributor Name": "distributor_name",
    "Estatus": "operational_status",
    "Meses No Uso": "months_no_use",
    "Alertas": "source_alert",
}


@dataclass
class InspectionResult:
    sha256: str
    extension: str
    media_type: str
    byte_size: int
    content_kind: str


@dataclass
class ParserOutput:
    parser_key: str
    parser_version: str
    headers: list[str]
    normalized_headers: list[str]
    rows: list[dict[str, Any]]
    row_count: int
    warnings: list[str]
    worksheets: list[str]
    sample_rows: list[dict[str, Any]]
    text_excerpt: str | None = None
    extraction_status: str = "extracted"


class FileInspectionService:
    def __init__(self, *, max_upload_size_bytes: int):
        self.max_upload_size_bytes = max_upload_size_bytes

    def inspect(self, *, filename: str, content: bytes, content_type: str | None) -> InspectionResult:
        if ".." in (filename or "") or "/" in (filename or "") or "\\" in (filename or ""):
            raise HTTPException(status_code=422, detail="invalid filename")
        size = len(content)
        if size <= 0:
            raise HTTPException(status_code=422, detail="empty file")
        if size > self.max_upload_size_bytes:
            raise HTTPException(status_code=422, detail="file exceeds max size")

        extension = Path(filename).suffix.lower()
        if extension in REJECTED_EXTENSIONS:
            raise HTTPException(status_code=422, detail="unsupported file type")
        if extension not in ALLOWED_EXTENSIONS:
            raise HTTPException(status_code=422, detail="file extension not allowed")

        media_type = self._detect_media_type(filename, content_type, content)
        content_kind = self._content_kind(extension, media_type)

        return InspectionResult(
            sha256=hashlib.sha256(content).hexdigest(),
            extension=extension,
            media_type=media_type,
            byte_size=size,
            content_kind=content_kind,
        )

    def _detect_media_type(self, filename: str, content_type: str | None, content: bytes) -> str:
        guessed, _ = mimetypes.guess_type(filename)
        if content.startswith(b"PK\x03\x04") and filename.lower().endswith(".xlsx"):
            return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if content.startswith(b"%PDF"):
            return "application/pdf"
        if content.startswith(b"\x89PNG"):
            return "image/png"
        if content.startswith(b"\xff\xd8\xff"):
            return "image/jpeg"
        return content_type or guessed or "application/octet-stream"

    def _content_kind(self, extension: str, media_type: str) -> str:
        if extension in {".csv", ".xlsx"}:
            return "structured_table"
        if extension == ".pdf":
            return "digital_pdf"
        if extension in {".png", ".jpg", ".jpeg"}:
            return "image"
        if extension == ".json":
            return "json"
        if extension == ".txt":
            return "text"
        return "unsupported"


class DocumentParserRegistry:
    version = "7.3A"

    def __init__(self, *, preview_row_limit: int):
        self.preview_row_limit = preview_row_limit

    def list_parsers(self) -> list[dict[str, Any]]:
        return [
            {"parser_key": "xlsx_table", "version": self.version, "supported_media_types": ["application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"]},
        ]

    def extract(self, *, document: DocumentRecord, filename: str, content: bytes, enable_ocr: bool = False) -> ParserOutput:
        ext = Path(filename).suffix.lower()
        if ext == ".xlsx":
            return self._parse_xlsx(content)
        raise HTTPException(status_code=422, detail="unsupported parser for file")

    def _decode_text(self, content: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    def _normalize_header(self, value: str) -> str:
        return re.sub(r"\s+", "_", value.strip().lower())

    def _parse_csv(self, content: bytes) -> ParserOutput:
        text = self._decode_text(content)
        sample = text[:2048]
        dialect = csv.Sniffer().sniff(sample, delimiters=",;") if sample.strip() else csv.excel
        reader = csv.DictReader(StringIO(text), dialect=dialect)
        headers = [item or "" for item in (reader.fieldnames or [])]
        rows: list[dict[str, Any]] = []
        for row in reader:
            if any((value or "").strip() for value in row.values()):
                rows.append({key: (value or "").strip() for key, value in row.items()})
        normalized = [self._normalize_header(item) for item in headers]
        return ParserOutput(
            parser_key="csv_table",
            parser_version=self.version,
            headers=headers,
            normalized_headers=normalized,
            rows=rows,
            row_count=len(rows),
            warnings=[f"delimiter:{getattr(dialect, 'delimiter', ',')}"] if headers else ["missing_headers"],
            worksheets=[],
            sample_rows=rows[: min(self.preview_row_limit, 20)],
        )

    def _parse_xlsx(self, content: bytes) -> ParserOutput:
        from io import BytesIO

        workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
        sheets = workbook.sheetnames
        weekly_headers = {"store_id", "TPV_HISTORICO", "ULTIMA_TX_DATE", "DIAS_SIN_TX"}
        inactive_headers = {"Store ID", "Meses No Uso", "Alertas", "Estatus"}
        candidates: list[tuple[int, str]] = []
        for sheet_name in sheets:
            first_row = next(workbook[sheet_name].iter_rows(values_only=True), ())
            headers = {str(value).strip() for value in first_row if value is not None}
            candidates.append((max(len(weekly_headers & headers), len(inactive_headers & headers)), sheet_name))

        score, selected_sheet_name = max(candidates, default=(0, sheets[0]))
        sheet = workbook[selected_sheet_name]
        rows_iter = sheet.iter_rows(values_only=True)
        first = next(rows_iter, None)
        headers = [str(item).strip() if item is not None else "" for item in (first or [])]
        normalized = [self._normalize_header(item) for item in headers]
        rows: list[dict[str, Any]] = []
        warnings: list[str] = [] if score >= 4 else ["recognized_report_structure_not_detected"]
        for values in rows_iter:
            raw = {headers[index]: values[index] for index in range(min(len(headers), len(values)))}
            if not any(value not in (None, "") for value in raw.values()):
                continue
            rows.append({key: value for key, value in raw.items()})
        return ParserOutput(
            parser_key="xlsx_table",
            parser_version=self.version,
            headers=headers,
            normalized_headers=normalized,
            rows=rows,
            row_count=len(rows),
            warnings=warnings,
            worksheets=[selected_sheet_name, *[name for name in sheets if name != selected_sheet_name]],
            sample_rows=[],
        )

    def _parse_pdf(self, content: bytes) -> ParserOutput:
        from io import BytesIO

        reader = PdfReader(BytesIO(content))
        pages = len(reader.pages)
        texts = [(page.extract_text() or "").strip() for page in reader.pages]
        full_text = "\n".join(item for item in texts if item)
        warnings: list[str] = []
        status = "extracted"
        if len(full_text.strip()) < 30:
            warnings.append("ocr_required")
            status = "ocr_required"
        return ParserOutput(
            parser_key="digital_pdf",
            parser_version=self.version,
            headers=[],
            normalized_headers=[],
            rows=[],
            row_count=0,
            warnings=[f"page_count:{pages}", *warnings],
            worksheets=[],
            sample_rows=[],
            text_excerpt=full_text[:1000] or None,
            extraction_status=status,
        )

    def _parse_text(self, content: bytes) -> ParserOutput:
        text = self._decode_text(content)
        return ParserOutput(
            parser_key="text",
            parser_version=self.version,
            headers=[],
            normalized_headers=[],
            rows=[],
            row_count=0,
            warnings=[],
            worksheets=[],
            sample_rows=[],
            text_excerpt=text[:2000],
        )

    def _parse_json(self, content: bytes) -> ParserOutput:
        data = json.loads(self._decode_text(content))
        if isinstance(data, list):
            rows = [item for item in data if isinstance(item, dict)]
            headers = list(rows[0].keys()) if rows else []
        elif isinstance(data, dict):
            rows = [data]
            headers = list(data.keys())
        else:
            rows = []
            headers = []
        return ParserOutput(
            parser_key="json",
            parser_version=self.version,
            headers=headers,
            normalized_headers=[self._normalize_header(item) for item in headers],
            rows=rows,
            row_count=len(rows),
            warnings=[] if rows else ["json_has_no_object_rows"],
            worksheets=[],
            sample_rows=rows[: min(len(rows), 20)],
            text_excerpt=None,
        )

    def _parse_image(self, content: bytes, *, enable_ocr: bool) -> ParserOutput:
        provider = DocumentIntelligenceProvider()
        result = provider.analyze_document(
            mime_type="image/png",
            extracted_text=None,
            enable_ocr=enable_ocr,
            binary=content,
        )
        status = "extracted" if result.text else "ocr_required"
        return ParserOutput(
            parser_key="image_registration",
            parser_version=self.version,
            headers=[],
            normalized_headers=[],
            rows=[],
            row_count=0,
            warnings=[result.method],
            worksheets=[],
            sample_rows=[],
            text_excerpt=result.text[:1000] if result.text else None,
            extraction_status=status,
        )


class DeterministicReportClassifier:
    REPORT_TYPES = [
        "netpay_weekly_sales_report",
        "netpay_inactive_stores_report",
        "unknown",
    ]

    def classify(self, *, filename: str, parser_output: ParserOutput) -> dict[str, Any]:
        headers = set(parser_output.headers)
        weekly_signals = {"store_id", "TPV_HISTORICO", "ULTIMA_TX_DATE", "DIAS_SIN_TX"}
        inactive_signals = {"Store ID", "Meses No Uso", "Alertas", "Estatus"}
        weekly_matched = sorted(weekly_signals & headers)
        inactive_matched = sorted(inactive_signals & headers)
        if len(weekly_matched) == len(weekly_signals):
            report_type = "netpay_weekly_sales_report"
            matched = weekly_matched
            missing_expected_fields = []
        elif len(inactive_matched) == len(inactive_signals):
            report_type = "netpay_inactive_stores_report"
            matched = inactive_matched
            missing_expected_fields = []
        else:
            report_type = "unknown"
            matched = []
            missing_expected_fields = sorted(weekly_signals - headers)
        confidence = 1.0 if report_type != "unknown" else 0.0

        return {
            "report_type": report_type,
            "confidence": confidence,
            "matched_signals": [f"header:{header}" for header in matched],
            "missing_expected_fields": missing_expected_fields,
        }


def propose_column_mappings(headers: list[str], report_type: str) -> list[dict[str, Any]]:
    mapping_catalog = {
        "netpay_weekly_sales_report": NETPAY_WEEKLY_SALES_MAPPING,
        "netpay_inactive_stores_report": NETPAY_INACTIVE_STORES_MAPPING,
    }.get(report_type, {})
    proposals: list[dict[str, Any]] = []
    for header in headers:
        best_key = mapping_catalog.get(header)
        best_confidence = 1.0 if best_key else 0.0
        proposals.append(
            {
                "source_header": header,
                "canonical_field": best_key,
                "confidence": best_confidence,
                "requires_review": best_confidence < 0.85,
            }
        )
    return proposals


def generate_candidate_observations(
    db: Session,
    *,
    source: SourceRecord,
    document: DocumentRecord,
    parser_output: ParserOutput,
    mappings: list[dict[str, Any]],
    report_type: str,
) -> tuple[int, dict[str, int | list[str]]]:
    mapping_by_header = {item["source_header"]: item for item in mappings if item.get("canonical_field")}
    count = 0

    existing = db.scalar(select(Observation).where(Observation.document_id == document.id).limit(1))
    if existing:
        observations = db.scalars(select(Observation).where(Observation.document_id == document.id)).all()
        return len(observations), _summary_from_observations(observations)

    if report_type == "netpay_inactive_stores_report":
        return _generate_inactive_store_observations(
            db,
            source=source,
            document=document,
            parser_output=parser_output,
            mappings=mappings,
        )

    stores: set[str] = set()
    churn_candidates = 0
    churn_insufficient_data = 0
    missing_identity = 0
    identity_conflicts = 0
    store_identity: dict[str, set[tuple[str, str]]] = {}
    for row_index, row in enumerate(parser_output.rows, start=1):
        facts = {
            mapping_by_header[header]["canonical_field"]: value
            for header, value in row.items()
            if header in mapping_by_header and value not in (None, "")
        }
        subject_reference = _reference(facts.get("store_id") or facts.get("client_id"))
        if subject_reference:
            stores.add(subject_reference)
        for header, value in row.items():
            mapping = mapping_by_header.get(header)
            if not mapping:
                continue
            if value in (None, ""):
                continue
            canonical = mapping["canonical_field"]
            observation = Observation(
                document_id=document.id,
                source_id=source.id,
                domain="netpay",
                subject_type="store" if canonical in {"store_id", "inactive_days", "sales_volume"} else "unknown",
                subject_reference=subject_reference or None,
                field_name=canonical,
                observed_value={"value": _json_value(value)},
                normalized_value={"value": _json_value(value)},
                identifier_type=canonical if canonical in {"store_id", "client_id"} else None,
                extraction_method=parser_output.parser_key,
                confidence=mapping["confidence"],
                confirmation_status="candidate",
                source_reference=f"row:{row_index};column:{header}",
                observed_at=utc_now(),
                provenance={
                    "worksheet": parser_output.worksheets[0] if parser_output.worksheets else None,
                    "row": row_index,
                    "column": header,
                    "parser_key": parser_output.parser_key,
                    "parser_version": parser_output.parser_version,
                    "mapping_confidence": mapping["confidence"],
                },
            )
            db.add(observation)
            count += 1

        inactive_days = _number(facts.get("inactive_days"))
        churn_eligible, insufficiency_reason = _churn_eligibility(facts)
        if subject_reference and churn_eligible:
            _add_finding(db, document, source, subject_reference, "finding.churn_candidate", row_index, parser_output, {"inactive_days": inactive_days})
            churn_candidates += 1
            count += 1
        elif subject_reference and insufficiency_reason:
            _add_finding(db, document, source, subject_reference, "finding.churn_insufficient_data", row_index, parser_output, {"reason": insufficiency_reason})
            churn_insufficient_data += 1
            count += 1
        if not subject_reference or not facts.get("client_id") or not facts.get("branch_name"):
            _add_finding(db, document, source, subject_reference or f"row:{row_index}", "finding.missing_identity", row_index, parser_output, {
                "missing": [name for name in ("store_id", "client_id", "branch_name") if not facts.get(name)],
            })
            missing_identity += 1
            count += 1
        if subject_reference:
            store_identity.setdefault(subject_reference, set()).add((_reference(facts.get("client_id")), _reference(facts.get("branch_name"))))

    for store_id, identities in store_identity.items():
        identities.discard(("", ""))
        if len(identities) > 1:
            _add_finding(db, document, source, store_id, "finding.identity_conflict", 0, parser_output, {"identity_versions": len(identities)})
            identity_conflicts += 1
            count += 1
    db.flush()
    return count, {
        "stores": len(stores),
        "churn_candidates": churn_candidates,
        "churn_insufficient_data": churn_insufficient_data,
        "critical_stores": 0,
        "assets_without_store": 0,
        "identity_conflicts": identity_conflicts,
        "missing_identity": missing_identity,
        "unavailable_fields": ["legal_company", "asset_serial", "asset_store_assignment", "inactivity_trend"],
    }


def _json_value(value: Any) -> Any:
    return value.isoformat() if isinstance(value, datetime) else value


def _generate_inactive_store_observations(
    db: Session,
    *,
    source: SourceRecord,
    document: DocumentRecord,
    parser_output: ParserOutput,
    mappings: list[dict[str, Any]],
) -> tuple[int, dict[str, int | list[str]]]:
    mapping_by_header = {item["source_header"]: item for item in mappings if item.get("canonical_field")}
    confirmed_store_ids = set(db.scalars(select(Observation.subject_reference).where(
        Observation.domain == "netpay",
        Observation.field_name == "store_id",
        Observation.confirmation_status == "confirmed",
    )).all())
    count = 0
    stores: set[str] = set()
    summary = {
        "stores": 0,
        "churn_watch": 0,
        "churn_candidates": 0,
        "churn_insufficient_data": 0,
        "cancellation_reviews": 0,
        "operational_blocks": 0,
        "activation_failures": 0,
        "linked_weekly_sales": 0,
        "critical_stores": 0,
        "assets_without_store": 0,
        "identity_conflicts": 0,
        "missing_identity": 0,
        "unavailable_fields": ["legal_company", "branch", "asset_serial", "asset_store_assignment"],
    }

    for row_index, row in enumerate(parser_output.rows, start=1):
        facts = {
            mapping_by_header[header]["canonical_field"]: value
            for header, value in row.items()
            if header in mapping_by_header and value not in (None, "")
        }
        store_id = _reference(facts.get("store_id"))
        if not store_id:
            summary["missing_identity"] += 1
            _add_finding(db, document, source, f"row:{row_index}", "finding.missing_identity", row_index, parser_output, {"missing": ["store_id"]})
            count += 1
            continue
        stores.add(store_id)

        for header, value in row.items():
            mapping = mapping_by_header.get(header)
            if not mapping or value in (None, ""):
                continue
            canonical = mapping["canonical_field"]
            db.add(Observation(
                document_id=document.id,
                source_id=source.id,
                domain="netpay",
                subject_type="store",
                subject_reference=store_id,
                field_name=canonical,
                observed_value={"value": _json_value(value)},
                normalized_value={"value": _json_value(value)},
                identifier_type="store_id" if canonical == "store_id" else None,
                extraction_method=parser_output.parser_key,
                confidence=mapping["confidence"],
                confirmation_status="candidate",
                source_reference=f"row:{row_index};column:{header}",
                observed_at=utc_now(),
                provenance={"worksheet": parser_output.worksheets[0], "row": row_index, "column": header, "parser_key": parser_output.parser_key, "parser_version": parser_output.parser_version},
            ))
            count += 1

        months_no_use = _number(facts.get("months_no_use"))
        if months_no_use is None:
            _add_finding(db, document, source, store_id, "finding.churn_insufficient_data", row_index, parser_output, {"reason": "months_no_use_unknown"})
            summary["churn_insufficient_data"] += 1
            count += 1
        elif months_no_use >= 2:
            _add_finding(db, document, source, store_id, "finding.churn_candidate", row_index, parser_output, {"churn_status": "churn_candidate", "months_no_use": months_no_use, "requires_human_approval": True})
            summary["churn_candidates"] += 1
            count += 1
        elif months_no_use == 1:
            _add_finding(db, document, source, store_id, "finding.churn_watch", row_index, parser_output, {"churn_status": "watch", "months_no_use": months_no_use})
            summary["churn_watch"] += 1
            count += 1

        status = _normalized_text(facts.get("operational_status"))
        alert = _normalized_text(facts.get("source_alert"))
        if "cobrar renta o cancelar" in alert:
            _add_finding(db, document, source, store_id, "finding.cancellation_review", row_index, parser_output, {"recommended_action": "cancellation_review", "requires_human_approval": True})
            summary["cancellation_reviews"] += 1
            count += 1
        if "bloqueado terminal incorrecta" in status:
            _add_finding(db, document, source, store_id, "finding.operational_block", row_index, parser_output, {"recommended_action": "investigate_terminal_assignment"})
            summary["operational_blocks"] += 1
            count += 1
        elif "bloqueado cambio terminal" in status:
            _add_finding(db, document, source, store_id, "finding.operational_block", row_index, parser_output, {"recommended_action": "investigate_terminal_replacement"})
            summary["operational_blocks"] += 1
            count += 1
        if "sin activacion" in alert:
            _add_finding(db, document, source, store_id, "finding.activation_failure", row_index, parser_output, {"recommended_action": "investigate_activation"})
            summary["activation_failures"] += 1
            count += 1
        if store_id in confirmed_store_ids:
            summary["linked_weekly_sales"] += 1

    summary["stores"] = len(stores)
    db.flush()
    return count, summary


def _normalized_text(value: Any) -> str:
    return "".join(char for char in unicodedata.normalize("NFD", _reference(value)).casefold() if unicodedata.category(char) != "Mn")


def _reference(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value: Any) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _churn_eligibility(facts: dict[str, Any]) -> tuple[bool, str | None]:
    inactive_days = _number(facts.get("inactive_days"))
    if inactive_days is None:
        return False, "inactive_days_unknown"
    if inactive_days < 60:
        return False, None

    last_activity_at = facts.get("last_activity_at")
    report_cutoff_at = facts.get("report_cutoff_at")
    if not isinstance(last_activity_at, datetime):
        return False, "last_activity_at_missing_or_malformed"
    if not isinstance(report_cutoff_at, datetime):
        return False, "report_cutoff_at_missing_or_malformed"

    calculated_inactive_days = (report_cutoff_at - last_activity_at).days
    if calculated_inactive_days < 0 or abs(calculated_inactive_days - inactive_days) > 1:
        return False, "inactive_days_not_supported_by_activity_dates"
    return True, None


def _add_finding(db: Session, document: DocumentRecord, source: SourceRecord, subject_reference: str, field_name: str, row_index: int, parser_output: ParserOutput, value: dict) -> None:
    db.add(Observation(
        document_id=document.id,
        source_id=source.id,
        domain="netpay",
        subject_type="store",
        subject_reference=subject_reference,
        field_name=field_name,
        observed_value=value,
        normalized_value=value,
        extraction_method=parser_output.parser_key,
        confidence=1.0,
        confirmation_status="candidate",
        source_reference=f"row:{row_index};rule:{field_name}",
        observed_at=utc_now(),
        provenance={"worksheet": parser_output.worksheets[0], "row": row_index, "parser_key": parser_output.parser_key, "parser_version": parser_output.parser_version},
    ))


def _summary_from_observations(observations: list[Observation]) -> dict[str, int | list[str]]:
    field_names = [item.field_name for item in observations]
    return {
        "stores": len({item.subject_reference for item in observations if item.subject_reference and item.subject_type == "store"}),
        "churn_candidates": field_names.count("finding.churn_candidate"),
        "churn_insufficient_data": field_names.count("finding.churn_insufficient_data"),
        "critical_stores": field_names.count("finding.critical_store"),
        "assets_without_store": field_names.count("finding.asset_without_store"),
        "identity_conflicts": field_names.count("finding.identity_conflict"),
        "missing_identity": field_names.count("finding.missing_identity"),
        "unavailable_fields": ["legal_company", "asset_serial", "asset_store_assignment", "inactivity_trend"],
    }


class DataIntakePreviewService:
    def __init__(
        self,
        *,
        repository: LocalDocumentRepository,
        parser_registry: DocumentParserRegistry,
        classifier: DeterministicReportClassifier,
    ):
        self.repository = repository
        self.parsers = parser_registry
        self.classifier = classifier

    def process_document(self, db: Session, document: DocumentRecord, source: SourceRecord, *, force_reprocess: bool = False) -> dict[str, Any]:
        if not document.storage_reference:
            raise HTTPException(status_code=422, detail="document has no storage reference")

        binary = self.repository.open(document.storage_reference).read()
        parser_output = self.parsers.extract(document=document, filename=document.original_filename or document.filename or "document", content=binary)
        classification = self.classifier.classify(filename=document.original_filename or document.filename or "document", parser_output=parser_output)
        mappings = propose_column_mappings(parser_output.headers, classification["report_type"])

        existing_obs = db.scalars(select(Observation).where(Observation.document_id == document.id)).all()
        if force_reprocess or not existing_obs:
            obs_count, operational_summary = generate_candidate_observations(
                db,
                source=source,
                document=document,
                parser_output=parser_output,
                mappings=mappings,
                report_type=classification["report_type"],
            )
        else:
            obs_count = len(existing_obs)
            operational_summary = _summary_from_observations(existing_obs)

        strong_identifiers = [item for item in mappings if item.get("canonical_field") in {"store_id", "client_id", "asset_serial", "rfc"}]
        conflicts = [item for item in mappings if item.get("requires_review")]

        preview = {
            "document_metadata": {
                "document_id": str(document.id),
                "filename": document.original_filename or document.filename,
                "storage_reference": document.storage_reference,
            },
            "duplicate_status": bool(document.duplicate_of_document_id),
            "detected_report_type": classification["report_type"],
            "confidence": classification["confidence"],
            "parser_used": parser_output.parser_key,
            "worksheets": parser_output.worksheets,
            "original_columns": parser_output.headers,
            "proposed_canonical_mappings": mappings,
            "sample_rows": parser_output.sample_rows,
            "total_rows": parser_output.row_count,
            "candidate_observation_count": obs_count,
            "detected_strong_identifiers": strong_identifiers,
            "conflicts": conflicts,
            "warnings": [*parser_output.warnings, *classification["missing_expected_fields"]],
            "fields_requiring_review": [item for item in mappings if item["requires_review"]],
            "proposed_changes_summary": {
                "new_observations": obs_count,
                "report_type": classification["report_type"],
            },
            "operational_summary": operational_summary,
        }

        document.detected_file_type = self._detected_file_type(document.media_type)
        document.detected_report_type = classification["report_type"]
        document.classification_confidence = classification["confidence"]
        document.parser_key = parser_output.parser_key
        document.parser_version = parser_output.parser_version
        document.row_count = parser_output.row_count
        document.observation_count = obs_count
        document.extraction_status = parser_output.extraction_status
        document.review_status = "preview_ready"
        document.processing_error = None
        document.processed_at = utc_now()
        meta = dict(document.metadata_json or {})
        meta["preview"] = preview
        meta["classification"] = classification
        meta["mappings"] = mappings
        document.metadata_json = meta

        record_event(
            db,
            event_type="data_intake.document_classified",
            aggregate_type="document_record",
            aggregate_id=document.id,
            payload={"report_type": classification["report_type"], "confidence": classification["confidence"]},
        )
        record_event(
            db,
            event_type="data_intake.extraction_completed",
            aggregate_type="document_record",
            aggregate_id=document.id,
            payload={"parser_key": parser_output.parser_key, "row_count": parser_output.row_count},
        )
        record_event(
            db,
            event_type="data_intake.preview_ready",
            aggregate_type="document_record",
            aggregate_id=document.id,
            payload={"observation_count": obs_count},
        )

        return preview

    def _detected_file_type(self, media_type: str) -> str:
        if "csv" in media_type or "spreadsheet" in media_type:
            return "structured_table"
        if media_type == "application/pdf":
            return "digital_pdf"
        if media_type.startswith("image/"):
            return "image"
        if media_type in {"application/json", "text/json"}:
            return "json"
        if media_type.startswith("text/"):
            return "text"
        return "unsupported"
