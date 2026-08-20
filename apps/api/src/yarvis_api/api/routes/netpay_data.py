"""D1 Netpay operational data: deterministic preview and human acceptance."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
import zipfile
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from uuid import UUID

from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, Query, Response, UploadFile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from yarvis_api.api.dependencies.authority import authority_envelope
from yarvis_api.application.authority import IdentityAuthorityEnvelope
from yarvis_api.application.netpay_inbox_authority import netpay_inbox_envelope
from yarvis_api.database import get_db
from yarvis_api.models.domain_event import record_event
from yarvis_api.models.netpay_master import NetpayClient, NetpayStoreReference
from yarvis_api.models.netpay_operational_data import (
    NoUsageCampaignEntry,
    OperationalDataBatch,
    OperationalDataCommandReceipt,
    OperationalDataRow,
    StoreProfitabilityFact,
)
from yarvis_api.schemas.netpay_operational_data import (
    OperationalDataAcceptance,
    OperationalDataBatchPage,
    OperationalDataBatchRead,
    OperationalDataRowRead,
    OperationalRowResolution,
)

router = APIRouter(prefix="/netpay/data", tags=["netpay-operational-data"])
_MAX_BYTES, _MAX_ROWS, _MAX_COLUMNS, _MAX_SHEETS = 5_242_880, 5_000, 50, 10
_DATASETS = {"monthly_store_profitability", "no_usage_campaign"}
_REQUIRED = {
    "monthly_store_profitability": (
        {"storeid"},
        {"mes", "periodo", "reportingperiod"},
        {"rentabilidad", "profitability"},
    ),
    "no_usage_campaign": (
        {"rfc", "rfcdistribuidor", "rfcdistributor"},
        {"storeid"},
        {"periodo", "campaignperiod", "mes"},
        {"mesessinuso", "monthswithoutuse"},
    ),
}


def _operator(envelope):
    netpay_inbox_envelope(envelope, required_scope="netpay.inbox.manage")
    return envelope


def _viewer(envelope):
    netpay_inbox_envelope(envelope, required_scope="netpay.inbox.read")
    return envelope


def _normal(value: object) -> str:
    return re.sub(
        r"[^a-z0-9]", "", unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().casefold()
    )


def _store_id(value: object) -> tuple[str | None, str | None]:
    if value is None or isinstance(value, bool):
        return None, "missing_store_id"
    if isinstance(value, float):
        if not value.is_integer():
            return None, "invalid_store_id"
        value = int(value)
    text = unicodedata.normalize("NFKC", str(value)).strip().lstrip("'").strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if not text:
        return None, "missing_store_id"
    if len(text) > 100 or not re.fullmatch(r"[A-Za-z0-9_-]+", text):
        return None, "invalid_store_id"
    return text, None


def _filename(name: str) -> str:
    value = Path(name or "dataset").name
    if not re.fullmatch(r"[A-Za-z0-9._ -]{1,255}", value) or Path(value).suffix.lower() not in {".xlsx", ".csv"}:
        raise HTTPException(422, "unsupported_file")
    return value


def _number(value: object) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return None


def _integer(value: object) -> int | None:
    number = _number(value)
    return int(number) if number is not None and number.is_integer() else None


def _period(value: object) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    match = re.fullmatch(r"\s*(20\d{2})[-/](0?[1-9]|1[0-2])\s*", str(value or ""))
    return f"{match.group(1)}-{int(match.group(2)):02d}" if match else None


def _find(row: dict[str, object], *names: str) -> object | None:
    values = {_normal(key): value for key, value in row.items()}
    return next((values[_normal(name)] for name in names if _normal(name) in values), None)


def _valid_headers(headers: list[str], dataset_type: str) -> bool:
    normalized = [_normal(header) for header in headers]
    return (
        bool(normalized)
        and all(normalized)
        and len(normalized) == len(set(normalized))
        and all(set(normalized) & group for group in _REQUIRED[dataset_type])
    )


def _read_rows(headers: list[str], values) -> list[dict[str, object]]:
    rows = []
    for raw in values:
        if any(isinstance(value, str) and value.startswith("=") for value in raw):
            raise HTTPException(422, "formulas_not_allowed")
        item = {headers[index]: raw[index] for index in range(min(len(headers), len(raw)))}
        if any(value not in (None, "") for value in item.values()):
            rows.append(item)
    return rows


def _parse(filename: str, content: bytes, dataset_type: str) -> tuple[list[dict[str, object]], str]:
    if not content or len(content) > _MAX_BYTES:
        raise HTTPException(422, "invalid_file_size")
    if Path(filename).suffix.lower() == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(422, "invalid_csv_encoding") from None
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;")
        except csv.Error:
            dialect = csv.excel
        values = csv.reader(StringIO(text), dialect=dialect)
        headers = [str(value).strip() for value in next(values, [])]
        if not _valid_headers(headers, dataset_type):
            raise HTTPException(422, "missing_or_ambiguous_headers")
        rows, sheet_name = _read_rows(headers, values), "CSV"
    else:
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                names = archive.namelist()
            if any(name.startswith("xl/externalLinks/") or name.endswith("vbaProject.bin") for name in names):
                raise HTTPException(422, "unsafe_workbook")
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        except (zipfile.BadZipFile, OSError, ValueError, KeyError):
            raise HTTPException(422, "corrupt_workbook") from None
        if not workbook.sheetnames or len(workbook.sheetnames) > _MAX_SHEETS:
            raise HTTPException(422, "invalid_worksheet_count")
        candidates = []
        for sheet in workbook.worksheets:
            headers = [
                str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(values_only=True), ())
            ]
            if _valid_headers(headers, dataset_type):
                candidates.append((sheet.title, headers))
        if len(candidates) != 1:
            raise HTTPException(422, "worksheet_missing_or_ambiguous")
        sheet_name, headers = candidates[0]
        values = workbook[sheet_name].iter_rows(values_only=True)
        next(values, None)
        rows = _read_rows(headers, values)
    if len(headers) > _MAX_COLUMNS or len(rows) > _MAX_ROWS:
        raise HTTPException(422, "dataset_limits_exceeded")
    return rows, sheet_name


def _digest(value: object) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"), default=str).encode()).hexdigest()


def _store_state(db: Session, organization_id: UUID, normalized_ids: set[str]) -> str:
    refs = (
        db.scalars(
            select(NetpayStoreReference).where(
                NetpayStoreReference.organization_id == organization_id,
                NetpayStoreReference.normalized_store_id.in_(normalized_ids),
                NetpayStoreReference.active.is_(True),
            )
        ).all()
        if normalized_ids
        else []
    )
    return _digest(
        [
            (str(item.id), item.normalized_store_id, item.updated_at.isoformat())
            for item in sorted(refs, key=lambda value: str(value.id))
        ]
    )


def _latest_fingerprint(
    db: Session, dataset_type: str, organization_id: UUID, store_reference_id: UUID, period: str
) -> str | None:
    model = StoreProfitabilityFact if dataset_type == "monthly_store_profitability" else NoUsageCampaignEntry
    period_column = model.reporting_period if dataset_type == "monthly_store_profitability" else model.campaign_period
    fact = db.scalar(
        select(model)
        .where(
            model.organization_id == organization_id,
            model.store_reference_id == store_reference_id,
            period_column == period,
        )
        .order_by(model.created_at.desc(), model.id.desc())
    )
    return fact.row_fingerprint if fact else None


def _row_read(row: OperationalDataRow) -> OperationalDataRowRead:
    payload, store = row.controlled_payload or {}, str((row.controlled_payload or {}).get("store_id") or "")
    return OperationalDataRowRead(
        id=row.id,
        source_row_number=row.source_row_number,
        validation_status=row.validation_status,
        match_status=row.match_status,
        store_reference_id=row.store_reference_id,
        error_codes=row.error_codes or [],
        projected_action=row.projected_action,
        preview={
            "store_id": f"***{store[-4:]}" if store else None,
            "reporting_period": payload.get("reporting_period"),
            "product_uen": payload.get("product_uen"),
            "months_without_usage": payload.get("months_without_usage"),
            "merchant_status": payload.get("merchant_status"),
            "alert_code": payload.get("alert_code"),
        },
    )


def _batch_read(
    db: Session, batch: OperationalDataBatch, include_rows: bool = True, duplicate: bool = False
) -> OperationalDataBatchRead:
    rows = (
        db.scalars(
            select(OperationalDataRow)
            .where(OperationalDataRow.batch_id == batch.id)
            .order_by(OperationalDataRow.source_row_number)
        ).all()
        if include_rows
        else []
    )
    return OperationalDataBatchRead.model_validate(batch).model_copy(
        update={
            "rows": [_row_read(row) for row in rows],
            "hash_identifier": batch.source_hash[:12],
            "duplicate_upload": duplicate,
        }
    )


def _event(db, envelope, batch, event_type):
    record_event(
        db,
        event_type=event_type,
        aggregate_type="netpay_operational_data_batch",
        aggregate_id=batch.id,
        organization_id=envelope.organization_id,
        payload={"dataset_type": batch.dataset_type, "status": batch.status, "row_counts": batch.row_counts},
    )


def _target(db, envelope, batch_id):
    batch = db.scalar(
        select(OperationalDataBatch).where(
            OperationalDataBatch.id == batch_id, OperationalDataBatch.organization_id == envelope.organization_id
        )
    )
    if not batch:
        raise HTTPException(404, "not_found")
    return batch


def _receipt(db, envelope, key, command, payload, mutation):
    fingerprint = _digest(payload)
    receipt = db.scalar(
        select(OperationalDataCommandReceipt).where(
            OperationalDataCommandReceipt.organization_id == envelope.organization_id,
            OperationalDataCommandReceipt.command_type == command,
            OperationalDataCommandReceipt.idempotency_key == key,
        )
    )
    if receipt:
        if receipt.request_fingerprint != fingerprint:
            raise HTTPException(409, "idempotency_conflict")
        return receipt.result_status_code, receipt.result_response_body
    batch, code = mutation()
    body = _batch_read(db, batch).model_dump(mode="json")
    db.add(
        OperationalDataCommandReceipt(
            organization_id=envelope.organization_id,
            command_type=command,
            idempotency_key=key,
            request_fingerprint=fingerprint,
            actor_principal_id=envelope.principal_id,
            result_batch_id=batch.id,
            result_status_code=code,
            result_response_body=body,
        )
    )
    return code, body


@router.post("/datasets", response_model=OperationalDataBatchRead, status_code=201)
async def upload_dataset(
    response: Response,
    file: UploadFile = File(...),
    dataset_type: str = Form(...),
    rfc_filter: str | None = Form(None),
    idempotency_key: str = Header(..., alias="Idempotency-Key", min_length=1, max_length=255),
    db: Session = Depends(get_db),
    envelope: IdentityAuthorityEnvelope = Depends(authority_envelope),
):
    current = _operator(envelope)
    if dataset_type not in _DATASETS:
        raise HTTPException(422, "unsupported_dataset")
    if dataset_type == "no_usage_campaign" and not rfc_filter:
        raise HTTPException(422, "rfc_filter_required")
    filename, content = _filename(file.filename or ""), await file.read()
    source_hash = hashlib.sha256(content).hexdigest()
    existing = db.scalar(
        select(OperationalDataBatch).where(
            OperationalDataBatch.organization_id == current.organization_id,
            OperationalDataBatch.dataset_type == dataset_type,
            OperationalDataBatch.source_hash == source_hash,
        )
    )
    if existing:
        response.status_code = 200
        return _batch_read(db, existing, duplicate=True)
    raw, selected_sheet = _parse(filename, content, dataset_type)
    authorized = raw
    if dataset_type == "no_usage_campaign":
        expected = _normal(rfc_filter)
        if not expected or any(_find(row, "rfc", "rfc distribuidor", "rfc distributor") is None for row in raw):
            raise HTTPException(422, "rfc_filter_unverifiable")
        authorized = [
            row for row in raw if _normal(_find(row, "rfc", "rfc distribuidor", "rfc distributor")) == expected
        ]

    def mutation():
        batch = OperationalDataBatch(
            organization_id=current.organization_id,
            dataset_type=dataset_type,
            reporting_period=None,
            selected_sheet=selected_sheet,
            sanitized_filename=filename,
            source_hash=source_hash,
            preview_token="0" * 64,
            store_state_hash="0" * 64,
            status="validating",
            row_counts={},
            uploaded_by_principal_id=current.principal_id,
        )
        db.add(batch)
        db.flush()
        candidates, natural_keys, normalized_ids = [], {}, set()
        for number, source in enumerate(authorized, start=2):
            store, store_error = _store_id(_find(source, "store id", "store_id"))
            period = _period(_find(source, "mes", "periodo", "reporting period", "campaign period"))
            errors = [item for item in (store_error, None if period else "missing_reporting_period") if item]
            profitability, months = (
                _number(_find(source, "rentabilidad", "profitability")),
                _integer(_find(source, "meses sin uso", "months without use")),
            )
            if dataset_type == "monthly_store_profitability" and profitability is None:
                errors.append("invalid_profitability")
            if dataset_type == "no_usage_campaign" and months is None:
                errors.append("invalid_months_without_usage")
            client_ref = str(_find(source, "client id", "client_id") or "").strip() or None
            if client_ref and not db.scalar(
                select(NetpayClient.id).where(
                    NetpayClient.organization_id == current.organization_id,
                    NetpayClient.external_reference == client_ref,
                )
            ):
                errors.append("client_reference_unmatched")
            payload = {
                "store_id": store,
                "client_external_reference": client_ref,
                "reporting_period": period,
                "product_uen": str(_find(source, "producto", "product", "uen") or "").strip() or None,
                "volume": _number(_find(source, "volumen", "volume")),
                "transaction_count": _integer(_find(source, "transacciones", "transaction count")),
                "income": _number(_find(source, "ingresos", "income")),
                "cost": _number(_find(source, "costos", "cost")),
                "commissions": _number(_find(source, "comisiones", "commissions")),
                "profitability": profitability,
                "months_without_usage": months,
                "merchant_status": str(_find(source, "estatus", "estado", "merchant status") or "").strip()[:64]
                or None,
                "alert_code": str(_find(source, "alerta", "alert") or "").strip()[:64] or None,
                "outcome_code": str(_find(source, "resultado", "outcome") or "").strip()[:64] or None,
                "follow_up_code": str(_find(source, "seguimiento", "follow up") or "").strip()[:64] or None,
            }
            row_key = _digest([_normal(store), period])
            if store and period:
                natural_keys[row_key] = natural_keys.get(row_key, 0) + 1
                normalized_ids.add(_normal(store))
            candidates.append((number, payload, errors, store, row_key))
        state_hash = _store_state(db, current.organization_id, normalized_ids)
        counts = {
            key: 0
            for key in (
                "valid",
                "invalid",
                "matched",
                "unmatched",
                "ambiguous",
                "duplicates",
                "conflicts",
                "projected_inserts",
                "projected_updates",
                "unchanged",
            )
        }
        periods, preview_material = set(), []
        for number, payload, errors, store, row_key in candidates:
            if natural_keys.get(row_key, 0) > 1:
                errors.append("duplicate_row")
                counts["duplicates"] += 1
            refs = (
                db.scalars(
                    select(NetpayStoreReference).where(
                        NetpayStoreReference.organization_id == current.organization_id,
                        NetpayStoreReference.normalized_store_id == _normal(store),
                        NetpayStoreReference.active.is_(True),
                    )
                ).all()
                if store
                else []
            )
            match = "matched" if len(refs) == 1 else "unmatched" if len(refs) == 0 else "ambiguous"
            counts[match] += 1
            validation = "invalid" if errors else "valid" if match == "matched" else "needs_review"
            counts["invalid" if validation == "invalid" else "valid"] += 1
            row_fingerprint = _digest(payload)
            projected = "invalid" if validation == "invalid" else "conflict" if match != "matched" else "insert"
            if projected == "conflict":
                counts["conflicts"] += 1
            elif projected == "insert":
                previous = _latest_fingerprint(
                    db, dataset_type, current.organization_id, refs[0].id, payload["reporting_period"]
                )
                projected = "unchanged" if previous == row_fingerprint else "update" if previous else "insert"
                counts[
                    {"insert": "projected_inserts", "update": "projected_updates", "unchanged": "unchanged"}[projected]
                ] += 1
            db.add(
                OperationalDataRow(
                    batch_id=batch.id,
                    source_row_number=number,
                    controlled_payload=payload,
                    validation_status=validation,
                    match_status=match,
                    store_reference_id=refs[0].id if len(refs) == 1 else None,
                    error_codes=errors,
                    row_key=row_key,
                    row_fingerprint=row_fingerprint,
                    projected_action=projected,
                )
            )
            preview_material.append([number, row_key, row_fingerprint, validation, match, projected])
            if payload["reporting_period"]:
                periods.add(payload["reporting_period"])
        batch.reporting_period = next(iter(periods)) if len(periods) == 1 else None
        batch.store_state_hash, batch.preview_token = (
            state_hash,
            _digest([dataset_type, source_hash, selected_sheet, state_hash, preview_material]),
        )
        batch.row_counts = {"received": len(raw), "authorized": len(authorized), **counts}
        batch.status = "needs_review"
        _event(db, current, batch, "netpay.operational_dataset.uploaded")
        _event(db, current, batch, "netpay.operational_dataset.validated")
        return batch, 201

    try:
        code, body = _receipt(
            db,
            current,
            idempotency_key,
            "IC-NETPAY-CMD-026",
            {"dataset_type": dataset_type, "source_hash": source_hash},
            mutation,
        )
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "concurrent_upload_conflict") from error
    response.status_code = code
    return body


@router.get("/datasets", response_model=OperationalDataBatchPage)
def list_datasets(
    dataset_type: str | None = Query(None),
    db: Session = Depends(get_db),
    envelope: IdentityAuthorityEnvelope = Depends(authority_envelope),
):
    current = _viewer(envelope)
    statement = select(OperationalDataBatch).where(OperationalDataBatch.organization_id == current.organization_id)
    if dataset_type:
        statement = statement.where(OperationalDataBatch.dataset_type == dataset_type)
    items = db.scalars(statement.order_by(OperationalDataBatch.created_at.desc(), OperationalDataBatch.id)).all()
    return OperationalDataBatchPage(items=[_batch_read(db, item, False) for item in items], total=len(items))


@router.get("/datasets/{batch_id}", response_model=OperationalDataBatchRead)
def get_dataset(
    batch_id: UUID, db: Session = Depends(get_db), envelope: IdentityAuthorityEnvelope = Depends(authority_envelope)
):
    current = _viewer(envelope)
    return _batch_read(db, _target(db, current, batch_id))


@router.post("/datasets/{batch_id}/accept", response_model=OperationalDataBatchRead)
def accept_dataset(
    batch_id: UUID,
    payload: OperationalDataAcceptance,
    response: Response,
    idempotency_key: str = Header(..., alias="Idempotency-Key", min_length=1, max_length=255),
    db: Session = Depends(get_db),
    envelope: IdentityAuthorityEnvelope = Depends(authority_envelope),
):
    current = _operator(envelope)

    def mutation():
        batch = _target(db, current, batch_id)
        if batch.status == "accepted":
            return batch, 200
        if batch.status != "needs_review" or payload.preview_token != batch.preview_token:
            raise HTTPException(409, "preview_expired_or_mismatched")
        rows = db.scalars(
            select(OperationalDataRow)
            .where(OperationalDataRow.batch_id == batch.id)
            .order_by(OperationalDataRow.source_row_number)
            .with_for_update()
        ).all()
        normalized = {
            _normal(row.controlled_payload.get("store_id")) for row in rows if row.controlled_payload.get("store_id")
        }
        if _store_state(db, current.organization_id, normalized) != batch.store_state_hash:
            raise HTTPException(409, "preview_expired_or_mismatched")
        if not rows or any(
            row.validation_status != "valid" or row.match_status != "matched" or row.store_reference_id is None
            for row in rows
        ):
            raise HTTPException(409, "dataset_requires_review")
        for row in rows:
            if row.projected_action == "unchanged":
                continue
            item = row.controlled_payload
            common = {
                "organization_id": current.organization_id,
                "batch_id": batch.id,
                "store_reference_id": row.store_reference_id,
                "row_fingerprint": row.row_fingerprint,
            }
            if batch.dataset_type == "monthly_store_profitability":
                db.add(
                    StoreProfitabilityFact(
                        **common,
                        client_external_reference=item["client_external_reference"],
                        reporting_period=item["reporting_period"],
                        product_uen=item["product_uen"],
                        volume=item["volume"],
                        transaction_count=item["transaction_count"],
                        income=item["income"],
                        cost=item["cost"],
                        commissions=item["commissions"],
                        profitability=item["profitability"],
                        no_use_indicator=None,
                    )
                )
            else:
                db.add(
                    NoUsageCampaignEntry(
                        **common,
                        campaign_period=item["reporting_period"],
                        months_without_usage=item["months_without_usage"],
                        merchant_status=item["merchant_status"],
                        alert_code=item["alert_code"],
                        outcome_code=item["outcome_code"],
                        follow_up_code=item["follow_up_code"],
                    )
                )
        batch.status = "accepted"
        _event(db, current, batch, "netpay.operational_dataset.accepted")
        return batch, 200

    try:
        code, body = _receipt(
            db,
            current,
            idempotency_key,
            "IC-NETPAY-CMD-028",
            {"batch_id": str(batch_id), "preview_token": payload.preview_token},
            mutation,
        )
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(409, "concurrent_confirmation_conflict") from error
    response.status_code = code
    return body


@router.post("/datasets/{batch_id}/reject", response_model=OperationalDataBatchRead)
def reject_dataset(
    batch_id: UUID,
    response: Response,
    idempotency_key: str = Header(..., alias="Idempotency-Key", min_length=1, max_length=255),
    db: Session = Depends(get_db),
    envelope: IdentityAuthorityEnvelope = Depends(authority_envelope),
):
    current = _operator(envelope)

    def mutation():
        batch = _target(db, current, batch_id)
        if batch.status == "accepted":
            raise HTTPException(409, "accepted_dataset_cannot_be_rejected")
        batch.status = "rejected"
        _event(db, current, batch, "netpay.operational_dataset.rejected")
        return batch, 200

    code, body = _receipt(db, current, idempotency_key, "IC-NETPAY-CMD-029", {"batch_id": str(batch_id)}, mutation)
    db.commit()
    response.status_code = code
    return body


@router.put("/datasets/{batch_id}/rows/{row_id}/match", response_model=OperationalDataBatchRead)
def resolve_match(
    batch_id: UUID,
    row_id: UUID,
    payload: OperationalRowResolution,
    response: Response,
    idempotency_key: str = Header(..., alias="Idempotency-Key", min_length=1, max_length=255),
    db: Session = Depends(get_db),
    envelope: IdentityAuthorityEnvelope = Depends(authority_envelope),
):
    current = _operator(envelope)

    def mutation():
        batch = _target(db, current, batch_id)
        row = db.scalar(
            select(OperationalDataRow).where(OperationalDataRow.id == row_id, OperationalDataRow.batch_id == batch.id)
        )
        reference = db.scalar(
            select(NetpayStoreReference).where(
                NetpayStoreReference.id == payload.store_reference_id,
                NetpayStoreReference.organization_id == current.organization_id,
                NetpayStoreReference.active.is_(True),
            )
        )
        if not row or not reference or batch.status != "needs_review":
            raise HTTPException(404, "not_found")
        row.store_reference_id, row.match_status, row.resolved_by_principal_id = (
            reference.id,
            "matched",
            current.principal_id,
        )
        row.validation_status = "valid" if not row.error_codes else "invalid"
        row.projected_action = "insert" if row.validation_status == "valid" else "invalid"
        rows = db.scalars(select(OperationalDataRow).where(OperationalDataRow.batch_id == batch.id)).all()
        batch.store_state_hash = _store_state(
            db,
            current.organization_id,
            {
                _normal(item.controlled_payload.get("store_id"))
                for item in rows
                if item.controlled_payload.get("store_id")
            },
        )
        batch.preview_token = _digest(
            [
                batch.source_hash,
                batch.store_state_hash,
                [(str(item.id), str(item.store_reference_id), item.validation_status) for item in rows],
            ]
        )
        counts = dict(batch.row_counts)
        counts.update(
            {
                "matched": sum(item.match_status == "matched" for item in rows),
                "unmatched": sum(item.match_status == "unmatched" for item in rows),
                "ambiguous": sum(item.match_status == "ambiguous" for item in rows),
            }
        )
        batch.row_counts = counts
        _event(db, current, batch, "netpay.operational_row_match.resolved")
        return batch, 200

    code, body = _receipt(
        db,
        current,
        idempotency_key,
        "IC-NETPAY-CMD-030",
        {"batch_id": str(batch_id), "row_id": str(row_id), "store_reference_id": str(payload.store_reference_id)},
        mutation,
    )
    db.commit()
    response.status_code = code
    return body


@router.get("/datasets/{batch_id}/results")
def dataset_results(
    batch_id: UUID, db: Session = Depends(get_db), envelope: IdentityAuthorityEnvelope = Depends(authority_envelope)
):
    current, batch = _viewer(envelope), None
    batch = _target(db, current, batch_id)
    if batch.status != "accepted":
        return []
    if batch.dataset_type == "monthly_store_profitability":
        rows = db.scalars(
            select(StoreProfitabilityFact).where(
                StoreProfitabilityFact.organization_id == current.organization_id,
                StoreProfitabilityFact.batch_id == batch.id,
            )
        ).all()
        return [
            {
                "store_reference_id": str(row.store_reference_id),
                "reporting_period": row.reporting_period,
                "product_uen": row.product_uen,
                "volume": row.volume,
                "transaction_count": row.transaction_count,
                "profitability": row.profitability,
                "result": "available",
            }
            for row in rows
        ]
    rows = db.scalars(
        select(NoUsageCampaignEntry).where(
            NoUsageCampaignEntry.organization_id == current.organization_id, NoUsageCampaignEntry.batch_id == batch.id
        )
    ).all()
    return [
        {
            "store_reference_id": str(row.store_reference_id),
            "campaign_period": row.campaign_period,
            "months_without_usage": row.months_without_usage,
            "merchant_status": row.merchant_status,
            "alert_code": row.alert_code,
            "result": "available",
        }
        for row in rows
    ]
